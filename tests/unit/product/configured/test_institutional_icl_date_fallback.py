from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl

from aip.product.configured.readers.institutional_icl_reader import InstitutionalICLReader


def _write_icl_workbook(path: Path, *, b7_value: object = None) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "ICL"
    worksheet["B7"] = b7_value

    rows = {
        10: (100000, 2.5, 2.0, 3.0),
        11: (200000, 200.0, 100.0, 90.0),
        12: (300000, 80.0, 50.0, 30.0),
        13: (310000, 100.0, 60.0, 40.0),
        14: (320000, 20.0, 10.0, 10.0),
    }
    for row_number, (code, total, mn, me) in rows.items():
        worksheet.cell(row=row_number, column=11, value=code)
        worksheet.cell(row=row_number, column=27, value=total)
        worksheet.cell(row=row_number, column=28, value=mn)
        worksheet.cell(row=row_number, column=29, value=me)

    workbook.save(path)


def test_reader_uses_institutional_filename_when_b7_is_empty(tmp_path: Path) -> None:
    path = tmp_path / "ICL 29 SETIEMBRE_2026.xlsx"
    _write_icl_workbook(path)

    result = InstitutionalICLReader().read(path)

    assert result.valuation_date == date(2026, 9, 29)
    assert result.diagnostics["source_cells"]["valuation_date"] == (
        "filename:ICL 29 SETIEMBRE_2026.xlsx"
    )
    assert any("institutional filename" in warning for warning in result.warnings)


def test_reader_keeps_b7_as_primary_date_source(tmp_path: Path) -> None:
    path = tmp_path / "ICL 29 SETIEMBRE_2026.xlsx"
    _write_icl_workbook(path, b7_value="30/09/2026")

    result = InstitutionalICLReader().read(path)

    assert result.valuation_date == date(2026, 9, 30)
    assert result.diagnostics["source_cells"]["valuation_date"] == "B7"
    assert not any("institutional filename" in warning for warning in result.warnings)
