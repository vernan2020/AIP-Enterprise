from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from aip.product.configured.irrbb.borrowing_workbook_inspector import (
    BorrowingWorkbookSchemaInspector,
)
from aip.product.configured.irrbb.physical_source_registry import BORROWING_WORKBOOK_SOURCE


def _governed_path(tmp_path: Path) -> Path:
    return tmp_path / BORROWING_WORKBOOK_SOURCE.logical_name


def _save_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Obligaciones"
    sheet.append(["Operacion", "Saldo", None, "Moneda"])
    sheet.append(["OP-001", 125000.0, None, "CRC"])
    sheet.append([None, None, None, None])
    sheet.append(["OP-002", 80000.0, None, "USD"])

    hidden = workbook.create_sheet("Parametros")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Version"
    hidden["B1"] = "2026"
    workbook.save(path)
    workbook.close()


def test_discovery_preserves_governed_identity_and_does_not_expose_parent_path(
    tmp_path: Path,
) -> None:
    path = _governed_path(tmp_path)
    _save_workbook(path)

    discovery = BorrowingWorkbookSchemaInspector().discover(path)
    expected_digest = hashlib.sha256(path.read_bytes()).hexdigest()

    assert discovery.source_id == BORROWING_WORKBOOK_SOURCE.source_id
    assert discovery.source_file_name == BORROWING_WORKBOOK_SOURCE.logical_name
    assert discovery.file_sha256 == expected_digest
    assert discovery.source_reference.endswith(f"#sha256={expected_digest}")
    assert str(tmp_path) not in discovery.source_reference
    assert "C:\\Users\\" not in discovery.source_reference


def test_discovery_reports_visible_and_hidden_sheet_topology_from_observed_values(
    tmp_path: Path,
) -> None:
    path = _governed_path(tmp_path)
    _save_workbook(path)

    discovery = BorrowingWorkbookSchemaInspector().discover(path)
    topology = {item.sheet_name: item for item in discovery.sheets}

    obligations = topology["Obligaciones"]
    assert obligations.visibility == "visible"
    assert obligations.observed_non_empty_row_count == 3
    assert obligations.observed_non_empty_cell_count == 9
    assert obligations.observed_last_non_empty_row == 4
    assert obligations.observed_last_non_empty_column == 4

    parameters = topology["Parametros"]
    assert parameters.visibility == "hidden"
    assert parameters.observed_non_empty_row_count == 1
    assert parameters.observed_non_empty_cell_count == 2


def test_declared_header_preserves_exact_labels_and_reports_blanks_and_duplicates(
    tmp_path: Path,
) -> None:
    path = _governed_path(tmp_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Obligaciones"
    sheet.append(["Operacion", " Tasa ", "   ", "operacion"])
    sheet.append(["OP-001", 0.07, None, "X"])
    workbook.save(path)
    workbook.close()

    inspection = BorrowingWorkbookSchemaInspector().inspect_declared_header(
        path,
        sheet_name="Obligaciones",
        header_row=1,
    )

    assert tuple(cell.column_letter for cell in inspection.cells) == ("A", "B", "C", "D")
    assert tuple(cell.label for cell in inspection.cells) == (
        "Operacion",
        " Tasa ",
        None,
        "operacion",
    )
    assert inspection.blank_column_indexes == (3,)
    assert inspection.duplicate_labels == ("Operacion",)


def test_declared_header_requires_explicit_existing_sheet_and_positive_row(tmp_path: Path) -> None:
    path = _governed_path(tmp_path)
    _save_workbook(path)
    inspector = BorrowingWorkbookSchemaInspector()

    with pytest.raises(ValueError, match="sheet_name is required"):
        inspector.inspect_declared_header(path, sheet_name=" ", header_row=1)

    with pytest.raises(ValueError, match="header_row must be positive"):
        inspector.inspect_declared_header(path, sheet_name="Obligaciones", header_row=0)

    with pytest.raises(KeyError, match="worksheet was not found"):
        inspector.inspect_declared_header(path, sheet_name="No existe", header_row=1)

    with pytest.raises(ValueError, match="header_row exceeds"):
        inspector.inspect_declared_header(path, sheet_name="Obligaciones", header_row=99)


def test_declared_header_rejects_non_text_values_instead_of_coercing_them(tmp_path: Path) -> None:
    path = _governed_path(tmp_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Obligaciones"
    sheet.append(["Operacion", 2026, "Saldo"])
    sheet.append(["OP-001", "Dato", 100.0])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="non-text value observed at B1"):
        BorrowingWorkbookSchemaInspector().inspect_declared_header(
            path,
            sheet_name="Obligaciones",
            header_row=1,
        )


def test_declared_header_rejects_row_without_text_labels(tmp_path: Path) -> None:
    path = _governed_path(tmp_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Obligaciones"
    sheet.append([" ", None, "   "])
    sheet.append(["Operacion", "Saldo", "Moneda"])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="declared header row contains no text labels"):
        BorrowingWorkbookSchemaInspector().inspect_declared_header(
            path,
            sheet_name="Obligaciones",
            header_row=1,
        )


def test_source_path_validation_is_fail_closed(tmp_path: Path) -> None:
    inspector = BorrowingWorkbookSchemaInspector()
    missing = _governed_path(tmp_path)

    with pytest.raises(FileNotFoundError, match="borrowing workbook was not found"):
        inspector.discover(missing)

    directory = tmp_path / BORROWING_WORKBOOK_SOURCE.logical_name
    directory.mkdir()
    with pytest.raises(ValueError, match="regular file"):
        inspector.discover(directory)
    directory.rmdir()

    wrong_suffix = tmp_path / "Auxiliar Obligaciones Entidades 2026.xlsm"
    wrong_suffix.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="must use the .xlsx format"):
        inspector.discover(wrong_suffix)

    wrong_name = tmp_path / "Otro Auxiliar.xlsx"
    wrong_name.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="does not match the governed physical source identity"):
        inspector.discover(wrong_name)


def test_corrupt_xlsx_is_rejected_without_fabricating_schema(tmp_path: Path) -> None:
    path = _governed_path(tmp_path)
    path.write_bytes(b"this is not an xlsx workbook")

    with pytest.raises(ValueError, match="borrowing workbook could not be read"):
        BorrowingWorkbookSchemaInspector().discover(path)
