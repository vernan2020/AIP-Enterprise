"""Qt smoke tests for read-only contractual reconciliation in the real Liquidez UI."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

from PySide6.QtCore import QDate
from PySide6.QtWidgets import QApplication, QPushButton, QTableWidget, QTabWidget

from aip.product.configured.liquidity.sugef210_scenario_comparison import (
    LiquidityFourScenarioComparison,
    LiquidityScenarioDifference,
)
from aip.ui.modules.liquidity.views.sugef210_projection_panel import SUGEF210ProjectionPanel


def test_contractual_tab_and_cutoff_invalidation_without_private_cash():
    app = QApplication.instance() or QApplication([])
    panel = SUGEF210ProjectionPanel()
    button = panel.findChild(QPushButton, "sugef210RunContractualReconciliation")
    table = panel.findChild(QTableWidget, "sugef210ContractualReconciliationDetail")
    tabs = panel.findChild(QTabWidget)
    assert app is not None
    assert button is not None and button.isEnabled()
    assert table is not None and tabs is not None
    cutoff = panel._cutoff
    panel._receive_contractual_audit(
        cutoff,
        {
            "status": "MATCHED_REFERENCE_ONLY",
            "cutoff": cutoff.isoformat(),
            "maturity_records": 20,
            "unique_maturity_rows_matched": 20,
            "principal_reconciliation": {
                "distinct_contracts_reconciled_exactly": 19,
                "negative_contracts_matching_signed_band_total": 1,
                "distinct_contracts_with_principal_difference": 0,
            },
            "contract_identities_without_maturity_bands_by_source": {
                "CAPTACIONES": 1,
                "BORROWING": 0,
            },
            "unbanded_date_evidence_by_sugef210_window": {
                "by_window": {"D01": 1},
                "next_day_concentration_review": {
                    "status": "NO_LARGE_NEXT_DAY_CLUSTER_DETECTED",
                    "next_day_unbanded_contracts": 1,
                    "future_dated_unbanded_contracts": 1,
                },
            },
        },
    )
    assert table.rowCount() == 22
    assert "SIN CERTIFICAR" in [table.item(row, 1).text() for row in range(table.rowCount())]
    assert "privado" not in table.item(0, 2).text().lower()
    assert tabs.currentWidget() is table

    next_cutoff = date(cutoff.year - 1, cutoff.month, cutoff.day)
    panel._on_date_changed(QDate(next_cutoff.year, next_cutoff.month, next_cutoff.day))
    assert table.rowCount() == 0
    panel.close()


def test_four_scenario_tab_renders_partial_delta_then_invalidates_cutoff(monkeypatch):
    from aip.ui.modules.liquidity.views import sugef210_projection_panel as module

    app = QApplication.instance() or QApplication([])
    panel = SUGEF210ProjectionPanel()
    table = panel.findChild(QTableWidget, "sugef210FourScenarioComparison")
    assert app is not None and table is not None
    assert table.rowCount() == 0

    row = LiquidityScenarioDifference(
        scenario="COMBINED",
        currency="CRC",
        band_key="D02",
        net_flow=Decimal("-20"),
        net_flow_delta_vs_base=Decimal("-20"),
        accumulated_flow_delta_vs_base=Decimal("-20"),
        projected_liquidity=None,
        projected_liquidity_delta_vs_base=None,
        financing_need=None,
        financing_need_delta_vs_base=None,
    )

    def comparison(_sessions):
        assert len(tuple(_sessions)) == 4
        return LiquidityFourScenarioComparison(
            cutoff=panel._cutoff.isoformat(),
            scenarios=("BASE", "IDIOSYNCRATIC", "SYSTEMIC", "COMBINED"),
            rows=(row,),
            missing_evidence_by_scenario=(),
        )

    monkeypatch.setattr(module, "compare_approved_sugef210_scenarios", comparison)
    panel._scenario_sessions.update(
        {"BASE": None, "IDIOSYNCRATIC": None, "SYSTEMIC": None, "COMBINED": None}
    )
    panel._refresh_four_scenario_comparison()
    assert table.rowCount() == 1
    assert table.item(0, 0).text() == "COMBINED"
    assert table.item(0, 3).text() == "-20"
    assert table.item(0, 4).text() == "N/D"
    assert table.item(0, 5).text() == "N/D"
    assert "NO CERTIFICADO" in table.item(0, 6).text()
    cutoff = panel._cutoff
    previous = date(cutoff.year - 1, cutoff.month, cutoff.day)
    panel._on_date_changed(QDate(previous.year, previous.month, previous.day))
    assert table.rowCount() == 0
    assert panel._scenario_sessions == {}
    panel.close()


def test_four_scenario_tab_blocks_mixed_source_snapshots(monkeypatch):
    from aip.ui.modules.liquidity.views import sugef210_projection_panel as module

    app = QApplication.instance() or QApplication([])
    panel = SUGEF210ProjectionPanel()
    table = panel.findChild(QTableWidget, "sugef210FourScenarioComparison")
    assert app is not None and table is not None

    def incompatible(_sessions):
        raise ValueError("Financial source snapshot does not match")

    monkeypatch.setattr(module, "compare_approved_sugef210_scenarios", incompatible)
    panel._scenario_sessions.update(
        {"BASE": None, "IDIOSYNCRATIC": None, "SYSTEMIC": None, "COMBINED": None}
    )
    panel._refresh_four_scenario_comparison()
    assert table.rowCount() == 1
    assert table.item(0, 0).text() == "BLOQUEADO"
    assert "no coinciden" in table.item(0, 6).text()
    panel.close()



def test_four_scenario_excel_export_preserves_numeric_deltas_and_unknown_cash(
    monkeypatch, tmp_path
):
    from openpyxl import load_workbook

    from aip.ui.modules.liquidity.views import sugef210_projection_panel as module

    app = QApplication.instance() or QApplication([])
    panel = SUGEF210ProjectionPanel()
    assert app is not None
    panel._report = {"cutoff": panel._cutoff.isoformat()}
    panel._scenario_sessions.update(
        {"BASE": None, "IDIOSYNCRATIC": None, "SYSTEMIC": None, "COMBINED": None}
    )
    row = LiquidityScenarioDifference(
        scenario="COMBINED",
        currency="CRC",
        band_key="D02",
        net_flow=Decimal("-30"),
        net_flow_delta_vs_base=Decimal("-30"),
        accumulated_flow_delta_vs_base=Decimal("-30"),
        projected_liquidity=None,
        projected_liquidity_delta_vs_base=None,
        financing_need=None,
        financing_need_delta_vs_base=None,
    )
    monkeypatch.setattr(
        module,
        "compare_approved_sugef210_scenarios",
        lambda _sessions: LiquidityFourScenarioComparison(
            cutoff=panel._cutoff.isoformat(),
            scenarios=("BASE", "IDIOSYNCRATIC", "SYSTEMIC", "COMBINED"),
            rows=(row,),
            missing_evidence_by_scenario=(
                ("COMBINED", ("UNAPPROVED_PERIMETER",)),
            ),
        ),
    )
    destination = tmp_path / "four-scenario-test.xlsx"
    monkeypatch.setattr(
        module.QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
    )
    panel.export_reconciliation()
    book = load_workbook(destination, read_only=True, data_only=True)
    try:
        sheet = book["Comparación 4 escenarios"]
        assert sheet.max_row == 2
        assert sheet["A2"].value == "COMBINED"
        assert sheet["F2"].value == -30
        assert sheet["I2"].value == "N/D"
        assert sheet["K2"].value == "N/D"
        assert sheet["L2"].value == "UNAPPROVED_PERIMETER"
        assert sheet["M2"].value == "NO CERTIFICADO"
    finally:
        book.close()
        panel.close()


def test_four_scenario_excel_export_blocks_missing_scenario_sessions(monkeypatch, tmp_path):
    from openpyxl import load_workbook

    from aip.ui.modules.liquidity.views import sugef210_projection_panel as module

    app = QApplication.instance() or QApplication([])
    panel = SUGEF210ProjectionPanel()
    assert app is not None
    panel._report = {"cutoff": panel._cutoff.isoformat()}
    destination = tmp_path / "four-scenario-incomplete.xlsx"
    monkeypatch.setattr(
        module.QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
    )
    panel.export_reconciliation()
    book = load_workbook(destination, read_only=True, data_only=True)
    try:
        sheet = book["Comparación 4 escenarios"]
        assert sheet.max_row == 2
        assert sheet["A2"].value == "INCOMPLETA"
        assert sheet["E2"].value == "N/D"
    finally:
        book.close()
        panel.close()
