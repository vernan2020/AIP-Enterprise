from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from aip.ui.modules.financial_analysis.viewmodels.financial_analysis_view_model import (
    FinancialMetricHistoryPointView,
    FinancialMetricHistorySeriesView,
)
from aip.ui.modules.financial_analysis.views.financial_history_panel import FinancialHistoryPanel
from aip.ui.modules.portfolio.models.portfolio_dashboard_point import PortfolioDashboardPoint
from aip.ui.modules.portfolio.models.portfolio_valuation_comparison import (
    PortfolioValuationChartPoint,
)
from aip.ui.modules.portfolio.views.portfolio_valuation_comparison_view import (
    PortfolioValuationComparisonView,
)
from aip.ui.modules.portfolio.views.portfolio_view import PortfolioView
from aip.ui.services.export_service import TableExportService


def test_financial_history_excel_preserves_unavailable_observation(tmp_path) -> None:
    history = FinancialMetricHistorySeriesView(
        code="ROA",
        label="ROA",
        unit="%",
        latest_value="0.00%",
        period_change="N/D",
        source_account="SUGEF 81000",
        available_points=2,
        total_points=3,
        points=(
            FinancialMetricHistoryPointView(
                iso_date="2026-07-31", date_label="jul-26", value=1.2, display_value="1.20%"
            ),
            FinancialMetricHistoryPointView(
                iso_date="2026-08-31", date_label="ago-26", value=None, display_value="N/D"
            ),
            FinancialMetricHistoryPointView(
                iso_date="2026-09-30", date_label="sep-26", value=0.0, display_value="0.00%"
            ),
        ),
    )
    book_path = TableExportService().export_workbook(
        tmp_path / "roa.xlsx",
        sheets=(FinancialHistoryPanel._excel_sheet(history),),
        metadata={"Entidad": "Entidad prueba", "Fuente": "SUGEF"},
    )
    book = load_workbook(book_path)
    sheet = book["ROA"]
    assert sheet["A2"].value.date() == date(2026, 7, 31)
    assert sheet["B2"].value == 1.2
    assert sheet["B3"].value is None
    assert sheet["C3"].value == "N/D"
    assert sheet["B4"].value == 0
    assert sheet["E4"].value == "SUGEF 81000"
    assert len(sheet._charts) == 1


def test_portfolio_rankings_export_native_editable_bar_chart(tmp_path) -> None:
    point = PortfolioDashboardPoint(
        label="BCCR", value=Decimal("33.7"), secondary_value=Decimal("100"), detail="CRC"
    )
    sheet = PortfolioView._chart_data_sheet("Concentración por emisor", (point,), "%")
    path = TableExportService().export_workbook(tmp_path / "emisor.xlsx", sheets=(sheet,))
    workbook = load_workbook(path)
    result = workbook["Concentración por emisor"]
    assert result["B2"].value == 33.7
    assert result["C2"].value == 100
    assert result._charts[0].__class__.__name__ == "BarChart"


def test_gain_loss_excel_keeps_signed_net_separate_from_gains_losses(tmp_path) -> None:
    point = PortfolioValuationChartPoint(
        label="Instrumento", value=Decimal("-20.5"), positive=Decimal("10"),
        negative=Decimal("-30.5"),
    )
    sheet = PortfolioValuationComparisonView._chart_sheet("G-P", (point,))
    path = TableExportService().export_workbook(tmp_path / "gp.xlsx", sheets=(sheet,))
    result = load_workbook(path)["G-P"]
    assert result["B2"].value == -20.5
    assert result["C2"].value == 10
    assert result["D2"].value == -30.5
    assert result._charts[0].__class__.__name__ == "BarChart"
