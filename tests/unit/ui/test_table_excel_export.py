from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QTableWidget, QTableWidgetItem

from aip.ui.widgets.table_excel_export import export_table_excel


def test_qt_table_excel_exports_only_visible_rows_and_numeric_user_roles(tmp_path) -> None:
    app = QApplication.instance() or QApplication([])
    assert app is not None
    table = QTableWidget(3, 3)
    table.setHorizontalHeaderLabels(["Serie", "VeR", "Oculta"])
    table.setItem(0, 0, QTableWidgetItem("ISIN-001"))
    var = QTableWidgetItem("₡123,456.78")
    var.setData(Qt.ItemDataRole.UserRole, Decimal("123456.78"))
    table.setItem(0, 1, var)
    table.setItem(1, 0, QTableWidgetItem("ISIN-002"))
    table.setItem(1, 1, QTableWidgetItem("0.12%"))
    table.setItem(2, 0, QTableWidgetItem("ISIN-003"))
    table.hideRow(1)
    table.hideColumn(2)

    count = export_table_excel(table, tmp_path / "riesgo.xlsx", visible_rows_only=True)
    assert count == 2
    book = load_workbook(tmp_path / "riesgo.xlsx")
    ws = book["Datos"]
    assert ws["A1"].value == "Serie"
    assert ws["B2"].value == 123456.78
    assert ws["A3"].value == "ISIN-003"
    assert ws["C1"].value is None
