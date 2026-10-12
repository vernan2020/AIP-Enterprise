from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook
from PySide6.QtWidgets import QTableWidget, QTableWidgetItem

from aip.ui.modules.financial_analysis.views.financial_analysis_view import FinancialAnalysisView
from aip.ui.modules.price_risk.models.price_risk_row import RiskChartPoint
from aip.ui.modules.price_risk.views.price_risk_view import PriceRiskView
from aip.ui.services.export_service import TableExportService


def test_risk_chart_excel_retains_decimal_without_recomputing(tmp_path) -> None:
    points = (
        RiskChartPoint(label="Título A", value=Decimal("0.155"), secondary_value=Decimal("0")),
        RiskChartPoint(label="Título B", value=Decimal("-0.025"), secondary_value=Decimal("0.75")),
    )
    sheets = (PriceRiskView._risk_chart_sheet("VeR por emisor", points, "%"),)
    path = TableExportService().export_workbook(tmp_path / "ver.xlsx", sheets=sheets)
    worksheet = load_workbook(path)["VeR por emisor"]
    assert worksheet["B2"].value == 0.155
    assert worksheet["B3"].value == -0.025
    assert worksheet["C2"].value == 0
    assert worksheet["C3"].value == 0.75
    assert worksheet._charts[0].__class__.__name__ == "BarChart"


def test_risk_pareto_exports_native_line_chart(tmp_path) -> None:
    points = (RiskChartPoint(label="Título A", value=Decimal("35.5")),)
    sheet = PriceRiskView._risk_chart_sheet("Pareto VeR", points, "%", pareto=True)
    path = TableExportService().export_workbook(tmp_path / "pareto.xlsx", sheets=(sheet,))
    assert load_workbook(path)["Pareto VeR"]._charts[0].__class__.__name__ == "LineChart"


def test_filtered_risk_table_excel_excludes_hidden_rows(qt_app) -> None:
    table = QTableWidget(2, 2)
    try:
        table.setHorizontalHeaderLabels(["Serie", "Contribución"])
        table.setItem(0, 0, QTableWidgetItem("Serie visible"))
        table.setItem(0, 1, QTableWidgetItem("0.15%"))
        table.setItem(1, 0, QTableWidgetItem("Serie oculta"))
        table.setItem(1, 1, QTableWidgetItem("0.35%"))
        table.setRowHidden(1, True)
        sheet = PriceRiskView._visible_table_sheet("Títulos VeR", table)
        assert sheet.headers == ("Serie", "Contribución")
        assert sheet.rows == (("Serie visible", "0.15%"),)
    finally:
        table.close()
        table.deleteLater()
        qt_app.processEvents()


def test_financial_table_export_honors_hidden_source_columns(qt_app) -> None:
    table = QTableWidget(2, 3)
    try:
        table.setHorizontalHeaderLabels(["Indicador", "Valor", "Cuenta fuente"])
        for row, label in enumerate(("ROA", "ROE")):
            table.setItem(row, 0, QTableWidgetItem(label))
            table.setItem(row, 1, QTableWidgetItem("1.00%"))
            table.setItem(row, 2, QTableWidgetItem("SUGEF"))
        table.setColumnHidden(2, True)
        table.setRowHidden(1, True)
        sheet = FinancialAnalysisView._visible_financial_table("Indicadores", table)
        assert sheet.headers == ("Indicador", "Valor")
        assert sheet.rows == (("ROA", "1.00%"),)
    finally:
        table.close()
        table.deleteLater()
        qt_app.processEvents()
