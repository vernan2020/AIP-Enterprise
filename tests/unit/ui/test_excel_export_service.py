from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from aip.ui.services.export_service import ExcelSheet, TableExportService


def test_export_records_creates_real_xlsx_preserving_zeros_and_missing(tmp_path) -> None:
    path = TableExportService().export_records(
        tmp_path / "reporte",
        headers=["Cuenta", "Saldo"],
        rows=[["Activo", Decimal("125.75")], ["Cero", 0], ["N/D", None]],
        export_format="excel",
    )
    wb = load_workbook(path)
    sheet = wb["Datos"]
    assert sheet["A1"].value == "Cuenta"
    assert sheet["B2"].value == 125.75
    assert sheet["B3"].value == 0
    assert sheet["B4"].value is None
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B4"


def test_export_workbook_adds_editable_chart_and_source_metadata(tmp_path) -> None:
    path = TableExportService().export_workbook(
        tmp_path / "historia.xlsx",
        metadata={
            "Fuente": "SUGEF",
            "Entidad": "Entidad de prueba",
            "N/D": "Dato no publicado; no equivale a cero",
        },
        sheets=(
            ExcelSheet(
                title="ROA histórico",
                headers=("Fecha", "ROA (%)", "Estado"),
                rows=(
                    (date(2026, 7, 31), 1.2, "Disponible"),
                    (date(2026, 8, 31), None, "N/D"),
                    (date(2026, 9, 30), 0.0, "Disponible"),
                ),
                chart_title="ROA histórico",
                unit="%",
            ),
        ),
    )
    book = load_workbook(path)
    assert book.sheetnames == ["Información", "ROA histórico"]
    assert book["Información"]["B2"].value == "SUGEF"
    sheet = book["ROA histórico"]
    assert sheet["A2"].value.date() == date(2026, 7, 31)
    assert sheet["B3"].value is None
    assert sheet["B4"].value == 0
    assert len(sheet._charts) == 1
    assert sheet._charts[0].__class__.__name__ == "LineChart"


def test_excel_text_does_not_become_executable_formulas(tmp_path) -> None:
    output = TableExportService().export_records(
        tmp_path / "seguridad.xlsx",
        headers=["Código"],
        rows=[["=HYPERLINK(\"https://example.com\")"], ["+SUM(1,2)"], ["@SUM(1,2)"]],
        export_format="excel",
    )
    sheet = load_workbook(output).active
    assert sheet is not None
    for index in (2, 3, 4):
        cell = sheet.cell(index, 1)
        assert cell.data_type != "f"
        assert cell.value.startswith("'")


def test_excel_rejects_inconsistent_rows_before_creation(tmp_path) -> None:
    path = tmp_path / "incorrecto.xlsx"
    with pytest.raises(ValueError, match="Inconsistent row width"):
        TableExportService().export_workbook(
            path,
            sheets=(ExcelSheet(title="Datos", headers=("Uno", "Dos"), rows=((1,),)),),
        )
    assert not path.exists()


def test_excel_sheet_titles_are_unique_and_valid(tmp_path) -> None:
    sheets = (
        ExcelSheet(title="Cuenta/Resultado", headers=("Valor",), rows=((1,),)),
        ExcelSheet(title="Cuenta/Resultado", headers=("Valor",), rows=((2,),)),
    )
    path = TableExportService().export_workbook(tmp_path / "nombres.xlsx", sheets=sheets)
    assert load_workbook(path).sheetnames == ["Cuenta_Resultado", "Cuenta_Resultado_2"]
