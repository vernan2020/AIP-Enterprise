from __future__ import annotations

from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QDate

from aip.domain.liquidity.cashflow.services.sugef210_projection import (
    LiquidityAssumption,
    LiquidityBalanceArea,
    LiquidityEvidence,
    LiquidityFlowSide,
    LiquidityProjectionFlow,
    project_sugef210,
)
from aip.domain.liquidity.cashflow.services.sugef210_scenario_rules import (
    ScenarioGeneratedEvidence,
)
from aip.product.configured.liquidity.sugef210_master_coupon_schedule import (
    ApprovedMasterCouponSchedule,
)
from aip.product.configured.liquidity.sugef210_master_maturity_bridge import (
    MasterPrincipalEvidence,
)
from aip.product.configured.liquidity.sugef210_opening_session import ConfiguredOpeningSession
from aip.product.configured.liquidity.sugef210_projection_assembler import (
    assemble_partial_institutional_projection,
)
from aip.ui.modules.liquidity.views.liquidity_view import LiquidityView
from aip.ui.modules.liquidity.views.sugef210_projection_panel import (
    SUGEF210ProjectionPanel,
)


def test_sugef210_panel_shows_compliant_horizons_without_demo_amounts(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        panel._date.setDate(QDate(2026, 9, 30))
        assert panel._windows.rowCount() == 17
        assert panel._windows.item(0, 4).text() == "N/D"
        assert panel._windows.item(0, 5).text() == "N/D"
        assert panel._windows.item(8, 1).text() == "08/10/2026"
        assert panel._windows.item(11, 2).text() == "30/10/2026"
        assert panel._windows.item(12, 2).text() == "29/11/2026"
        assert panel._windows.item(13, 2).text() == "29/12/2026"
        assert panel._windows.item(16, 2).text() == "30/09/2027"
        assert panel._accounts.rowCount() == 0
    finally:
        panel.close()


def test_sugef210_xml_inventory_visible_but_not_available_cash(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        panel._date.setDate(QDate(2026, 9, 30))
        panel._receive(
            panel._cutoff,
            {
                "status": "ACCOUNT_LEVEL_EVIDENCE_ONLY_NO_AVAILABLE_CASH",
                "xml_records_read": 3,
                "accounts": [
                    {
                        "account_code": "100",
                        "currency_source_code": "1",
                        "record_count": 3,
                        "source_principal": "150000",
                        "classification": "NOT_APPROVED_FOR_CURRENT_ACCOUNT",
                        "master_isin_matches": 1,
                        "unverified_master_overlap": 2,
                    }
                ],
            },
        )
        assert panel._accounts.rowCount() == 1
        assert panel._accounts.item(0, 3).text() == "150000"
        assert "NO CERTIFICADA" in panel._status.text()
        assert all(panel._windows.item(row, 4).text() == "N/D" for row in range(17))
        panel._date.setDate(QDate(2026, 10, 19))
        assert panel._accounts.rowCount() == 0
        assert "No se reutilizan" in panel._status.text()
    finally:
        panel.close()


def test_existing_liquidity_page_includes_sugef210_tab(qt_app):
    view = LiquidityView()
    try:
        names = [view._tabs.tabText(i) for i in range(view._tabs.count())]
        assert "Flujo proyectado · SUGEF 2-10" in names
        assert view._projection_sugef210._windows.rowCount() == 17
    finally:
        view.close()


def test_contractual_xml_source_status_not_a_payment_schedule(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        panel._date.setDate(QDate(2026, 9, 30))
        panel._receive(
            panel._cutoff,
            {
                "status": "ACCOUNT_LEVEL_EVIDENCE_ONLY_NO_AVAILABLE_CASH",
                "xml_records_read": 0,
                "accounts": [],
                "record_evidence": [],
                "contractual_source_inventory": [
                    {
                        "source_key": "CREDIT",
                        "file": "NEC2024_Operaciones_5103.Nuevo.xml",
                        "records": 2,
                        "source_currency_balances_not_cashflows": {"1": "100.00"},
                        "dates": [
                            {"field": "FechaVencimiento", "valid": 1, "missing": 1, "invalid": 0}
                        ],
                        "status": "SOURCE_FIELDS_INSPECTED_NOT_A_PAYMENT_SCHEDULE",
                    },
                    {
                        "source_key": "CAPTACIONES",
                        "file": None,
                        "status": "SOURCE_UNAVAILABLE_NOT_ASSUMED_ZERO",
                    },
                ],
            },
        )
        assert panel._xml_perimeters.rowCount() == 2
        assert "NO flujos" in panel._XML_PERIMETER_HEADERS[3]
        assert panel._xml_perimeters.item(0, 3).text() == "1: 100.00"
        assert panel._xml_perimeters.item(1, 2).text() == "N/D"
        assert panel._windows.item(0, 4).text() == "N/D"
    finally:
        panel.close()


def test_governed_projection_updates_only_verified_currency_and_resets_on_scenario(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        flow = LiquidityProjectionFlow(
            reference="CONTRACT-CREDIT-TEST-001",
            source="TEST-ONLY",
            currency="CRC",
            payment_date=cutoff + timedelta(days=31),
            amount=Decimal("125"),
            side=LiquidityFlowSide.INFLOW,
            area=LiquidityBalanceArea.ASSET,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        result = project_sugef210(
            cutoff=cutoff,
            scenario="BASE",
            flows=(flow,),
            opening_liquidity={"CRC": Decimal("10")},
            source_areas_covered=(LiquidityBalanceArea.ASSET,),
            opening_certification_references={"CRC": "ACCOUNT-RECON-TEST"},
        )
        assert not result.certified_comprehensive
        panel.present_governed_projection(result)
        assert panel._windows.item(12, 4).text() == "125"
        assert panel._governed_flow_details.rowCount() == 34
        assert panel._governed_flow_details.item(12, 2).text() == "125"
        assert panel._governed_flow_details.item(12, 5).text() == "N/D"
        assert panel._governed_flow_details.item(17 + 12, 2).text() == "N/D"
        assert panel._windows.item(13, 4).text() == "N/D"
        assert panel._windows.item(12, 5).text() == "N/D"
        assert panel._windows.item(12, 9).text() == "135"
        assert panel._windows.item(12, 10).text() == "N/D"
        assert panel._projected_kpis["CRC"].text() == "135"
        assert panel._projected_kpis["USD"].text() == "N/D"
        assert "PARCIAL" in panel._windows.item(12, 6).text()
        panel._scenario.setCurrentText("SYSTEMIC")
        assert panel._governed_projection is None
        assert panel._windows.item(12, 4).text() == "N/D"
        assert panel._governed_flow_details.rowCount() == 0
        with pytest.raises(ValueError, match="cutoff/scenario"):
            panel.present_governed_projection(result)
    finally:
        panel.close()


def test_partial_xml_cashflow_excel_masks_unverified_other_source_zeros(
    qt_app, tmp_path, monkeypatch
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        credit = LiquidityProjectionFlow(
            reference="TEST-ORIGINAL-CREDIT-SETTLEMENT",
            source="TEST-CONTRACTUAL-CREDIT-DOCUMENT",
            currency="CRC",
            payment_date=cutoff + timedelta(days=31),
            amount=Decimal("125"),
            side=LiquidityFlowSide.INFLOW,
            area=LiquidityBalanceArea.ASSET,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        projection = project_sugef210(
            cutoff=cutoff,
            scenario="BASE",
            flows=(credit,),
            opening_liquidity={"CRC": Decimal("10")},
        )
        panel.present_governed_projection(projection)
        assert panel._governed_flow_details.item(12, 2).text() == "125"
        assert panel._governed_flow_details.item(12, 5).text() == "N/D"
        assert panel._governed_flow_details.item(12, 7).text() == "N/D"
        destination = tmp_path / "only_approved_sources.xlsx"
        monkeypatch.setattr(
            "aip.ui.modules.liquidity.views.sugef210_projection_panel."
            "QFileDialog.getSaveFileName",
            lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
        )
        panel.export_reconciliation()
        assert destination.is_file()
        workbook = load_workbook(destination, read_only=True, data_only=True)
        try:
            sheet = workbook["Flujo proyectado 365d"]
            # CRC starts at Excel row 2, then 31-60d at zero-based band index 12.
            row = 2 + 12
            assert sheet.cell(row, 7).value == 125
            assert sheet.cell(row, 7).data_type == "n"
            assert sheet.cell(row, 8).value == "N/D"
            assert sheet.cell(row, 10).value == "N/D"
            assert sheet.cell(row, 12).value == "N/D"
            assert sheet.cell(row, 15).value == 125
            assert sheet.cell(row, 15).data_type == "n"
            assert sheet.cell(row, 17).value == 10
            assert sheet.cell(row, 18).value == 135
            assert sheet.cell(row, 20).value == "PARTIAL_EVIDENCE_NOT_CERTIFIED"
        finally:
            workbook.close()
    finally:
        panel.close()


def test_export_preserves_negative_gap_and_funding_need_as_numeric_cells(
    qt_app, tmp_path, monkeypatch
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        outflow = LiquidityProjectionFlow(
            reference="TEST-DATED-USD-DEPOSIT-OUTFLOW",
            source="TEST-CONTRACTUAL-DEPOSIT-EVIDENCE",
            currency="USD",
            payment_date=cutoff + timedelta(days=2),
            amount=Decimal("17.25"),
            side=LiquidityFlowSide.OUTFLOW,
            area=LiquidityBalanceArea.LIABILITY,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        projection = project_sugef210(
            cutoff=cutoff,
            scenario="BASE",
            flows=(outflow,),
            opening_liquidity={"USD": Decimal("10.00")},
        )
        panel.present_governed_projection(projection)
        destination = tmp_path / "numeric_gap.xlsx"
        monkeypatch.setattr(
            "aip.ui.modules.liquidity.views.sugef210_projection_panel."
            "QFileDialog.getSaveFileName",
            lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
        )
        panel.export_reconciliation()
        workbook = load_workbook(destination, read_only=True, data_only=True)
        try:
            sheet = workbook["Flujo proyectado 365d"]
            # CRC occupies the first 17 bands, USD the subsequent 17.
            row = 2 + 17 + 2
            for column, amount in (
                (10, 17.25),
                (15, -17.25),
                (16, -17.25),
                (17, 10),
                (18, -7.25),
                (19, 7.25),
            ):
                cell = sheet.cell(row, column)
                assert cell.value == amount
                assert cell.data_type == "n"
            assert sheet.cell(row, 7).value == "N/D"
            assert sheet.cell(row, 20).value == "PARTIAL_EVIDENCE_NOT_CERTIFIED"
        finally:
            workbook.close()
    finally:
        panel.close()


def test_governed_partial_projection_can_export_without_xml_inventory(
    qt_app, tmp_path, monkeypatch
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        result = project_sugef210(
            cutoff=cutoff,
            scenario="BASE",
            flows=(),
            opening_liquidity=None,
        )
        panel.present_governed_projection(result)
        destination = tmp_path / "sugef210_only_partially_supported.xlsx"
        monkeypatch.setattr(
            "aip.ui.modules.liquidity.views.sugef210_projection_panel."
            "QFileDialog.getSaveFileName",
            lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
        )
        panel.export_reconciliation()
        assert destination.is_file()
        workbook = load_workbook(destination, read_only=True, data_only=True)
        try:
            assert "Flujo proyectado 365d" in workbook.sheetnames
            assert "Captaciones vista XML210" in workbook.sheetnames
            assert workbook["Captaciones vista XML210"].max_row == 1
            projection = workbook["Flujo proyectado 365d"]
            assert projection.max_row == 35
            assert projection.cell(2, 6).value == "CRC"
            assert projection.cell(2, 15).value == "N/D"
            assert projection.cell(2, 7).value == "N/D"
            metadata = dict(workbook["Control"].iter_rows(values_only=True))
            assert metadata["flujo_integral_certificado"] == "NO"
            assert metadata["escenario"] == "BASE"
        finally:
            workbook.close()
    finally:
        panel.close()


def test_governed_projection_rejects_incomplete_horizon_and_currency_rows(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        correct = project_sugef210(cutoff=cutoff, scenario="BASE", flows=(), opening_liquidity=None)
        with pytest.raises(ValueError, match="fixed SUGEF210"):
            panel.present_governed_projection(replace(correct, windows=correct.windows[:-1]))
        with pytest.raises(ValueError, match="do not reconcile"):
            panel.present_governed_projection(replace(correct, rows=correct.rows[:-1]))
        assert panel._windows.item(12, 4).text() == "N/D"
    finally:
        panel.close()


def test_ui_loads_governed_opening_as_partial_waterfall_and_does_not_reuse_cutoff(
    qt_app, tmp_path, monkeypatch
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        assembly = assemble_partial_institutional_projection(
            cutoff=cutoff,
            scenario="BASE",
            opening_by_currency={"CRC": Decimal("1000"), "USD": Decimal("85.25")},
            opening_certifications={"CRC": "TEST-ICL-CRC", "USD": "TEST-ICL-USD"},
            opening_nonduplication_reference="TEST-ICL-SIGHT-VALIDATION",
        )
        session = ConfiguredOpeningSession(
            cutoff=cutoff,
            scenario="BASE",
            assembly=assembly,
            sight_records_approved=1,
            sight_source_sha256="a" * 64,
            bank_support_sha256="b" * 64,
            bank_document_sha256="c" * 64,
            demand_deposit_accounts_reviewed=1,
            demand_deposit_runoff_events_approved=2,
            demand_deposit_source_sha256="e" * 64,
            demand_deposit_native_reconciliations=(
                ("TEST-210", "SIGHT", "2", 1, Decimal("100.00")),
            ),
        )
        panel._receive_governed_opening(cutoff, "BASE", session)
        assert panel._governed_projection is assembly.projection
        assert panel._demand_deposit_reconciliation.rowCount() == 1
        assert panel._demand_deposit_reconciliation.item(0, 0).text() == "TEST-210"
        assert panel._demand_deposit_reconciliation.item(0, 2).text() == "2"
        assert panel._demand_deposit_reconciliation.item(0, 4).text() == "100.00"
        assert "PARCIAL" in panel._demand_deposit_reconciliation.item(0, 5).text()
        destination = tmp_path / "reconciliation_XML210.xlsx"
        monkeypatch.setattr(
            "aip.ui.modules.liquidity.views.sugef210_projection_panel."
            "QFileDialog.getSaveFileName",
            lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
        )
        panel.export_reconciliation()
        assert destination.is_file()
        workbook = load_workbook(destination, read_only=True, data_only=True)
        try:
            assert "Captaciones vista XML210" in workbook.sheetnames
            sheet = workbook["Captaciones vista XML210"]
            assert sheet.max_row == 2
            assert sheet.cell(2, 1).value == "TEST-210"
            assert sheet.cell(2, 3).value == "2"
            assert sheet.cell(2, 5).value == "100.00"
            assert "PARCIAL" in sheet.cell(2, 6).value
        finally:
            workbook.close()
        assert panel._projected_kpis["CRC"].text() == "1000"
        assert panel._projected_kpis["USD"].text() == "85.25"
        assert "Proyección PARCIAL" in panel._opening_status.text()
        assert panel._windows.item(12, 4).text() == "N/D"
        panel._date.setDate(QDate(2026, 10, 31))
        assert panel._governed_projection is None
        assert panel._projected_kpis["USD"].text() == "N/D"
        assert panel._demand_deposit_reconciliation.rowCount() == 0
        panel._receive_governed_opening(cutoff, "BASE", session)
        assert panel._governed_projection is None
    finally:
        panel.close()


def test_opening_UI_error_never_books_current_accounts(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        panel._date.setDate(QDate(2026, 9, 30))
        panel._governed_opening_error(panel._cutoff, "BASE", "ValueError")
        assert "BLOCKED_GOVERNED_OPENING" in panel._opening_status.text()
        assert panel._windows.item(0, 9).text() == "N/D"
    finally:
        panel.close()


def test_governed_master_cashflow_displays_contract_date_without_inferred_coupons(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        master_event = LiquidityProjectionFlow(
            reference="MASTER:" + "c" * 64 + ":ROW-3:PRINCIPAL",
            source="TEST-MASTER-ORIGINAL-CELL",
            currency="CRC",
            payment_date=cutoff + timedelta(days=31),
            amount=Decimal("120"),
            side=LiquidityFlowSide.INFLOW,
            area=LiquidityBalanceArea.ASSET,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        admitted = MasterPrincipalEvidence(
            cutoff=cutoff,
            master_sha256="c" * 64,
            positions_read=1,
            flows=(master_event,),
            excluded=(),
            principal_by_currency=(("CRC", Decimal("120")),),
        )
        assembly = assemble_partial_institutional_projection(
            cutoff=cutoff,
            scenario="BASE",
            master=admitted,
            master_maturity_recognition_reference="TEST-CUSTODIAN-SETTLEMENT",
            opening_by_currency={"CRC": Decimal("1000"), "USD": Decimal("85.25")},
            opening_certifications={"CRC": "TEST-CRC", "USD": "TEST-USD"},
            opening_nonduplication_reference="TEST-ICL-RECON",
        )
        session = ConfiguredOpeningSession(
            cutoff=cutoff,
            scenario="BASE",
            assembly=assembly,
            sight_records_approved=1,
            sight_source_sha256="a" * 64,
            bank_support_sha256="b" * 64,
            bank_document_sha256="d" * 64,
            master_principal_events_approved=1,
        )
        panel._receive_governed_opening(cutoff, "BASE", session)
        assert panel._windows.item(12, 4).text() == "120"
        assert panel._windows.item(12, 9).text() == "1120"
        assert panel._windows.item(12, 5).text() == "N/D"
        assert panel._projected_kpis["CRC"].text() == "1120"
        assert "1 vencimientos de principal" in panel._opening_status.text()
        assert "PARCIAL" in panel._windows.item(12, 6).text()
    finally:
        panel.close()


def test_offbalance_approved_systemic_stress_visible_as_usd_outflow_not_cash(qt_app):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        panel._scenario.setCurrentText("SYSTEMIC")
        policy_id = "SUGEF210:SYSTEMIC:TEST-OFFBALANCE:V1"
        policy = LiquidityAssumption(
            assumption_id=policy_id,
            rationale="TEST-ONLY-APPROVED-CONTINGENCY",
            approval_reference="TEST-ALCO-OFFBALANCE",
            valid_from=cutoff,
            valid_to=date(2026, 12, 31),
            source_reference="TEST-SIGNED-OFF-CONTRACT-REGISTRY",
            scenario="SYSTEMIC",
        )
        drawdown = LiquidityProjectionFlow(
            reference="BEHAVIORAL:" + policy_id + ":TEST-UNDRAWN-LINE",
            source="TEST-ONLY-REGISTER-SHA:NOT-CASH",
            currency="USD",
            payment_date=cutoff + timedelta(days=7),
            amount=Decimal("35"),
            side=LiquidityFlowSide.OUTFLOW,
            area=LiquidityBalanceArea.OFF_BALANCE,
            evidence=LiquidityEvidence.GOVERNED_ASSUMPTION,
            assumption_id=policy_id,
        )
        scenario = ScenarioGeneratedEvidence(
            cutoff=cutoff,
            scenario="SYSTEMIC",
            flows=(drawdown,),
            assumptions=(policy,),
            matched_exposure_count=1,
        )
        assembly = assemble_partial_institutional_projection(
            cutoff=cutoff,
            scenario="SYSTEMIC",
            approved_behavioral=scenario,
            opening_by_currency={"CRC": Decimal("1000"), "USD": Decimal("85.25")},
            opening_certifications={"CRC": "TEST-BANK-CRC", "USD": "TEST-BANK-USD"},
            opening_nonduplication_reference="TEST-NO-ICL-OVERLAP",
        )
        session = ConfiguredOpeningSession(
            cutoff=cutoff,
            scenario="SYSTEMIC",
            assembly=assembly,
            sight_records_approved=1,
            sight_source_sha256="a" * 64,
            bank_support_sha256="b" * 64,
            bank_document_sha256="c" * 64,
            contingent_scenario_events_approved=1,
            offbalance_commitments_reviewed=1,
        )
        panel._receive_governed_opening(cutoff, "SYSTEMIC", session)
        assert panel._governed_flow_details.rowCount() == 34
        assert panel._governed_flow_details.item(17 + 7, 7).text() == "35"
        assert panel._governed_flow_details.item(7, 7).text() == "N/D"
        assert panel._windows.item(7, 10).text() == "50.25"
        assert panel._projected_kpis["USD"].text() == "50.25"
        assert "SYSTEMIC aprobado" in panel._opening_status.text()
        assert "PARCIAL" in panel._windows.item(7, 6).text()
        panel._scenario.setCurrentText("BASE")
        assert panel._governed_projection is None
        assert panel._governed_flow_details.rowCount() == 0
    finally:
        panel.close()


def test_investment_reconciliation_separates_sight_principal_coupon_and_excel_numeric(
    qt_app, tmp_path, monkeypatch
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        principal = LiquidityProjectionFlow(
            reference="MASTER:" + "c" * 64 + ":ROW-1:PRINCIPAL",
            source="TEST-MASTER-ROW-1",
            currency="CRC",
            payment_date=cutoff + timedelta(days=31),
            amount=Decimal("120.00"),
            side=LiquidityFlowSide.INFLOW,
            area=LiquidityBalanceArea.ASSET,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        coupon = LiquidityProjectionFlow(
            reference="MASTER_COUPON:" + "c" * 64 + ":TEST-USD-CUPON",
            source="TEST-CUSTODY-USD",
            currency="USD",
            payment_date=cutoff + timedelta(days=31),
            amount=Decimal("5.25"),
            side=LiquidityFlowSide.INFLOW,
            area=LiquidityBalanceArea.ASSET,
            evidence=LiquidityEvidence.CONTRACTUAL,
        )
        master = MasterPrincipalEvidence(
            cutoff=cutoff,
            master_sha256="c" * 64,
            positions_read=1,
            flows=(principal,),
            excluded=(),
            principal_by_currency=(("CRC", Decimal("120.00")), ("USD", Decimal(0))),
        )
        coupons = ApprovedMasterCouponSchedule(
            cutoff=cutoff,
            events=(coupon,),
            master_sha256="c" * 64,
            ledger_sha256="e" * 64,
            custodian_sha256="f" * 64,
            nonduplication_references=frozenset({"TEST-COUPON-RECON"}),
        )
        assembly = assemble_partial_institutional_projection(
            cutoff=cutoff,
            scenario="BASE",
            master=master,
            master_maturity_recognition_reference="TEST-PRINCIPAL-SETTLEMENT",
            approved_investment_coupons=coupons,
            opening_by_currency={"CRC": Decimal("1000"), "USD": Decimal("85.25")},
            opening_certifications={"CRC": "TEST-ICL-CRC", "USD": "TEST-ICL-USD"},
            opening_nonduplication_reference="TEST-ICL-OPENING",
        )
        session = ConfiguredOpeningSession(
            cutoff=cutoff,
            scenario="BASE",
            assembly=assembly,
            sight_records_approved=1,
            sight_source_sha256="a" * 64,
            bank_support_sha256="b" * 64,
            bank_document_sha256="d" * 64,
            master_principal_events_approved=1,
            master_coupon_events_approved=1,
            master_source_sha256="c" * 64,
            master_coupon_custodian_sha256="f" * 64,
            sight_added_by_currency=(("CRC", Decimal("0")), ("USD", Decimal("75.25"))),
        )
        panel._receive_governed_opening(cutoff, "BASE", session)
        table = panel._investment_reconciliation
        assert table.rowCount() == 6
        assert table.item(1, 0).text() == "MASTER_RECOGNIZED_PRINCIPAL"
        assert table.item(1, 3).text() == "1"
        assert table.item(1, 5).text() == "120.00"
        assert table.item(3, 1).text().startswith("APERTURA")
        assert table.item(3, 5).text() == "75.25"
        assert table.item(3, 3).text() == "N/D"
        assert table.item(5, 0).text() == "APPROVED_MASTER_COUPONS"
        assert table.item(5, 5).text() == "5.25"
        assert table.item(5, 7).text() == "f" * 64
        destination = tmp_path / "investment_sources.xlsx"
        monkeypatch.setattr(
            "aip.ui.modules.liquidity.views.sugef210_projection_panel."
            "QFileDialog.getSaveFileName",
            lambda *_args, **_kwargs: (str(destination), "Excel (*.xlsx)"),
        )
        panel.export_reconciliation()
        workbook = load_workbook(destination, read_only=True, data_only=True)
        try:
            sheet = workbook["Conciliacion inversiones"]
            assert sheet.max_row == 7
            assert sheet.cell(3, 4).value == 1
            assert sheet.cell(3, 4).data_type == "n"
            assert sheet.cell(3, 6).value == 120
            assert sheet.cell(3, 6).data_type == "n"
            assert sheet.cell(5, 6).value == 75.25
            assert sheet.cell(5, 6).data_type == "n"
            assert sheet.cell(5, 4).value == "N/D"
            assert sheet.cell(7, 6).value == 5.25
            assert sheet.cell(7, 6).data_type == "n"
            assert sheet.cell(7, 8).value == "f" * 64
            assert dict(workbook["Control"].iter_rows(values_only=True))[
                "flujo_integral_certificado"
            ] == "NO"
        finally:
            workbook.close()
        panel._scenario.setCurrentText("SYSTEMIC")
        assert table.rowCount() == 0
        assert panel._governed_projection is None
    finally:
        panel.close()


def test_investment_reconciliation_missing_evidence_stays_nd_and_inconsistent_counts_block(
    qt_app
):
    panel = SUGEF210ProjectionPanel()
    try:
        cutoff = date(2026, 9, 30)
        panel._date.setDate(QDate(2026, 9, 30))
        assembly = assemble_partial_institutional_projection(
            cutoff=cutoff,
            scenario="BASE",
            opening_by_currency={"CRC": Decimal(20), "USD": Decimal(10)},
            opening_certifications={"CRC": "TEST-CRC", "USD": "TEST-USD"},
            opening_nonduplication_reference="TEST-OPENING-RECON",
        )
        session = ConfiguredOpeningSession(
            cutoff=cutoff,
            scenario="BASE",
            assembly=assembly,
            sight_records_approved=0,
            sight_source_sha256="a" * 64,
            bank_support_sha256="b" * 64,
            bank_document_sha256="d" * 64,
        )
        panel._receive_governed_opening(cutoff, "BASE", session)
        table = panel._investment_reconciliation
        assert table.rowCount() == 6
        assert all(table.item(i, 5).text() == "N/D" for i in range(6))
        assert all(table.item(i, 3).text() == "N/D" for i in range(6) if i % 3 == 0)
        panel._invalidate_governed_projection()
        with pytest.raises(ValueError, match="Investment event counts differ"):
            panel._receive_governed_opening(
                cutoff,
                "BASE",
                replace(session, master_coupon_events_approved=1),
            )
        assert table.rowCount() == 0
        assert panel._governed_projection is None
    finally:
        panel.close()

