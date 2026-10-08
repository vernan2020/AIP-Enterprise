from __future__ import annotations

import csv
import json
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Mapping, Sequence


@dataclass(frozen=True, slots=True)
class ExcelSheet:
    """Typed data, optionally accompanied by an editable Excel line chart."""

    title: str
    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    chart_title: str | None = None
    category_column: int = 1
    value_column: int = 2
    unit: str = ""


def _safe_text(value: str) -> str:
    """Prevent spreadsheet formula execution in untrusted source labels."""

    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal, bool, date, datetime)):
        return value
    return _safe_text(str(value))


def _worksheet_name(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", "_", name).strip().strip("'") or "Datos"
    base = clean[:31]
    result = base
    counter = 2
    while result.casefold() in used:
        suffix = f"_{counter}"
        result = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(result.casefold())
    return result


class TableExportService:
    """Centralized CSV/JSON/XLSX exports with no invented source values."""

    def export_records(
        self, path: Path | str, *, headers: list[str], rows: list[list[Any]], export_format: str
    ) -> str:
        target = Path(path)
        if export_format == "csv":
            target = target.with_suffix(".csv")
            with target.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(rows)
        elif export_format == "json":
            target = target.with_suffix(".json")
            payload = [{header: value for header, value in zip(headers, row)} for row in rows]
            target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        elif export_format == "excel":
            return self.export_workbook(
                target,
                sheets=(
                    ExcelSheet(
                        title="Datos",
                        headers=tuple(headers),
                        rows=tuple(tuple(row) for row in rows),
                    ),
                ),
            )
        else:
            raise ValueError(f"Unsupported export format: {export_format}")
        return str(target)

    def export_workbook(
        self,
        path: Path | str,
        *,
        sheets: Sequence[ExcelSheet],
        metadata: Mapping[str, str] | None = None,
    ) -> str:
        """Create a real XLSX with editable series and traceability.

        Missing numeric observations remain blank (not zero). Metadata makes
        that interpretation explicit, preserving published SUGEF semantics.
        """

        if not sheets:
            raise ValueError("At least one Excel sheet is required")
        for sheet in sheets:
            if not sheet.headers:
                raise ValueError("Excel sheet headers must not be empty")
            if any(len(row) != len(sheet.headers) for row in sheet.rows):
                raise ValueError(f"Inconsistent row width for {sheet.title}")

        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        target = Path(path).with_suffix(".xlsx")
        workbook = Workbook()
        workbook.remove(workbook.active)
        names: set[str] = set()

        if metadata:
            info = workbook.create_sheet(_worksheet_name("Información", names))
            info.append(("Campo", "Valor"))
            for key, value in metadata.items():
                info.append((_safe_text(str(key)), _safe_text(str(value))))
            info.column_dimensions["A"].width = 28
            info.column_dimensions["B"].width = 78
            info.freeze_panes = "A2"
            info["A1"].font = info["B1"].font = Font(bold=True, color="FFFFFF")
            info["A1"].fill = info["B1"].fill = PatternFill("solid", fgColor="005EB8")

        for specification in sheets:
            worksheet = workbook.create_sheet(_worksheet_name(specification.title, names))
            worksheet.append(tuple(_safe_text(str(h)) for h in specification.headers))
            for row in specification.rows:
                worksheet.append(tuple(_cell_value(value) for value in row))
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="005EB8")
                cell.alignment = Alignment(wrap_text=True)
            for column_index, header in enumerate(specification.headers, start=1):
                max_width = max(
                    len(str(header)),
                    *(len(str(row[column_index - 1])) for row in specification.rows),
                )
                worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                    max(max_width + 3, 14), 45
                )
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (date, datetime)):
                        cell.number_format = "dd/mm/yyyy"
                    elif isinstance(cell.value, (Decimal, float)):
                        cell.number_format = "#,##0.00;[Red](#,##0.00);0.00"

            if specification.chart_title and specification.rows:
                if not (1 <= specification.category_column <= len(specification.headers)):
                    raise ValueError("Invalid chart category column")
                if not (1 <= specification.value_column <= len(specification.headers)):
                    raise ValueError("Invalid chart value column")
                chart = LineChart()
                chart.title = specification.chart_title
                chart.y_axis.title = specification.unit
                chart.style = 13
                chart.height = 9
                chart.width = 17
                values = Reference(
                    worksheet,
                    min_col=specification.value_column,
                    min_row=1,
                    max_row=worksheet.max_row,
                )
                categories = Reference(
                    worksheet,
                    min_col=specification.category_column,
                    min_row=2,
                    max_row=worksheet.max_row,
                )
                chart.add_data(values, titles_from_data=True)
                chart.set_categories(categories)
                worksheet.add_chart(
                    chart, f"{get_column_letter(max(len(specification.headers) + 2, 6))}3"
                )

        # Avoid publishing a partially written XLSX if a save fails.
        temporary_path = ""
        try:
            with tempfile.NamedTemporaryFile(
                prefix=".aip_excel_", suffix=".xlsx", dir=target.parent, delete=False
            ) as handle:
                temporary_path = handle.name
            workbook.save(temporary_path)
            os.replace(temporary_path, target)
        finally:
            if temporary_path and os.path.exists(temporary_path):
                os.unlink(temporary_path)
        return str(target)
