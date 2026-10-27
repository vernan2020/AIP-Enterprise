"""Configured-only SUGEF 2-10 liquidity planning panel.

Reads institutional XML investment source on demand. It NEVER shows example
numbers and does not pretend incomplete source coverage is a complete forecast.
"""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

from PySide6.QtCore import QDate, Qt, QThread, Signal
from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QDateEdit,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from aip.domain.liquidity.cashflow.services.sugef210_projection import (
    LiquidityBalanceArea,
    LiquidityFlowSide,
    LiquidityProjectionResult,
    sugef210_windows,
)
from aip.product.configured.liquidity.financial_plan_source import (
    FinancialPlanSourceAudit,
    read_financial_plan_assumptions,
)
from aip.product.configured.liquidity.sugef210_opening_session import (
    ConfiguredOpeningSession,
    run_configured_sight_opening,
)
from aip.product.configured.liquidity.sugef210_operational_readiness import (
    build_operational_gates,
    inspect_obligation_identity_with_config,
)
from aip.product.configured.liquidity.sugef210_projection_assembler import (
    LiquidityAuditedCashEvent,
)
from aip.product.configured.liquidity.sugef210_reconciliation_summary import (
    summarize_approved_contractual_source_coverage,
    summarize_contractual_reconciliation,
)
from aip.product.configured.liquidity.sugef210_scenario_comparison import (
    compare_approved_sugef210_scenarios,
)
from aip.product.configured.liquidity.sugef210_source_audit_cli import (
    inspect_xml_investment_source,
)


def last_completed_calendar_month_end(today: date) -> date:
    """Suggest a completed monthly cutoff; does not claim source publication."""
    return today.replace(day=1) - timedelta(days=1)


class _SUGEF210SourceWorker(QThread):
    completed = Signal(object)
    failed = Signal(str)

    def __init__(self, cutoff: date) -> None:
        super().__init__()
        self._cutoff = cutoff

    def run(self) -> None:
        try:
            report = inspect_xml_investment_source(self._cutoff)
        except Exception as exc:  # UI must fail closed and never leak paths/tokens
            self.failed.emit(type(exc).__name__)
            return
        self.completed.emit(report)


class _SUGEF210ContractualWorker(QThread):
    completed = Signal(object)
    failed = Signal(str)

    def __init__(self, cutoff: date) -> None:
        super().__init__()
        self._cutoff = cutoff

    def run(self) -> None:
        try:
            audit = inspect_obligation_identity_with_config(self._cutoff)
        except Exception as exc:
            # The UI must not expose operation IDs, source files or balances.
            self.failed.emit(type(exc).__name__)
            return
        self.completed.emit(audit)


class _SUGEF210FinancialPlanWorker(QThread):
    completed = Signal(object)
    failed = Signal(str)

    def __init__(self, cutoff: date, source_path: str) -> None:
        super().__init__()
        self._cutoff = cutoff
        self._source_path = source_path

    def run(self) -> None:
        try:
            evidence = read_financial_plan_assumptions(self._source_path, cutoff=self._cutoff)
        except Exception as exc:
            self.failed.emit(type(exc).__name__)
            return
        self.completed.emit(evidence)


class _SUGEF210OpeningWorker(QThread):
    completed = Signal(object)
    failed = Signal(str)

    def __init__(self, cutoff: date, scenario: str, approval_path: str, bank_path: str) -> None:
        super().__init__()
        self._cutoff = cutoff
        self._scenario = scenario
        self._approval_path = approval_path
        self._bank_path = bank_path

    def run(self) -> None:
        try:
            result = run_configured_sight_opening(
                cutoff=self._cutoff,
                scenario=self._scenario,
                approval_manifest=self._approval_path,
                bank_support_manifest=self._bank_path,
            )
        except Exception as exc:
            self.failed.emit(type(exc).__name__)
            return
        self.completed.emit(result)


class SUGEF210ProjectionPanel(QWidget):
    """Liquidity horizon framework + REAL source coverage, no DEMO waterfall."""

    _XML_PERIMETER_HEADERS = (
        "Fuente",
        "Archivo",
        "Registros",
        "Saldos por moneda de origen (NO flujos)",
        "Fechas contractuales presentes / ausentes / inválidas",
        "Estado de calidad",
    )
    _PLAN_HEADERS = (
        "Mes",
        "Supuesto / Driver",
        "Valor 2026 (¢ millones ref.)",
        "Área",
        "Sentido",
        "Tratamiento",
        "Celda fuente",
    )
    _PLAN_PARAMETER_HEADERS = (
        "ID",
        "Parámetro / actividad",
        "Valor",
        "Unidad",
        "Periodicidad",
        "Tipo",
        "Ref. Plan",
        "Tratamiento Liquidez",
        "Celda fuente",
    )
    _WINDOW_HEADERS = (
        "Horizonte",
        "Desde",
        "Hasta",
        "Apertura",
        "CRC neto",
        "USD neto",
        "Estado",
        "Principal Maestro CRC*",
        "Principal Maestro USD*",
        "Saldo proyectado CRC†",
        "Saldo proyectado USD†",
    )
    _SOURCE_READINESS_HEADERS = ("Control", "Estado", "Acción siguiente")
    _EVENT_AUDIT_HEADERS = (
        "Grupo fuente",
        "Referencia evento",
        "Referencia origen",
        "Fecha pago",
        "Moneda original",
        "Importe nativo",
        "Área",
        "Sentido",
        "Evidencia",
        "Supuesto",
        "Banda SUGEF 2-10",
        "Horizonte",
        "Cobertura",
    )
    _EVENT_AUDIT_UI_LIMIT = 300
    _INVESTMENT_RECONCILIATION_HEADERS = (
        "Origen de inversión",
        "Naturaleza",
        "Moneda original",
        "Eventos a 365 días",
        "Eventos >365 días",
        "Monto nativo a 365 días / apertura",
        "Estado de cobertura",
        "SHA256 respaldo",
    )
    _GOVERNED_FLOW_DETAIL_HEADERS = (
        "Banda",
        "Moneda",
        "Cobro activos",
        "Pago activos",
        "Entrada pasivos",
        "Pago pasivos",
        "Entradas contingentes",
        "Salidas contingentes",
        "Entradas negocio",
        "Salidas negocio",
        "Neto parcial",
        "GAP acumulado parcial",
        "Saldo parcial",
        "Eventos documentados",
        "Certificación",
    )
    _DETAIL_HEADERS = (
        "Referencia XML",
        "ID original",
        "Cuenta",
        "Moneda origen",
        "Saldo principal",
        "Id instrumento",
        "ISIN",
        "Serie",
        "Conciliación con Maestro",
        "Clasificación",
    )
    _SOURCE_HEADERS = (
        "Cuenta XML",
        "Moneda XML",
        "Registros",
        "Principal XML",
        "Clasificación",
        "Coincide ISIN",
        "Sin conciliación",
    )

    def __init__(self) -> None:
        super().__init__()
        self.setObjectName("sugef210LiquidityProjectionPanel")
        self._worker: _SUGEF210SourceWorker | None = None
        self._contractual_worker: _SUGEF210ContractualWorker | None = None
        self._report: dict[str, object] | None = None
        self._plan_worker: _SUGEF210FinancialPlanWorker | None = None
        self._opening_worker: _SUGEF210OpeningWorker | None = None
        self._plan_audit: FinancialPlanSourceAudit | None = None
        self._governed_projection: LiquidityProjectionResult | None = None
        self._scenario_sessions: dict[str, ConfiguredOpeningSession] = {}
        self._audited_events: tuple[LiquidityAuditedCashEvent, ...] = ()
        self._cutoff = last_completed_calendar_month_end(date.today())
        main = QVBoxLayout(self)
        main.setContentsMargins(10, 10, 10, 10)
        main.setSpacing(8)

        heading = QLabel("FLUJO PROYECTADO DE FONDOS · SUGEF 2-10")
        heading.setStyleSheet("font-size:14pt; font-weight:700; color:#064F87;")
        main.addWidget(heading)
        explanation = QLabel(
            "Horizontes diarios 0–7, semanales 8–30, mensuales y trimestrales "
            "hasta un año. Activos, pasivos, fuera de balance y eventos de negocio; "
            "la proyección requiere flujos contractuales, saldos iniciales y "
            "supuestos institucionales aprobados. Sin datos DEMO."
        )
        explanation.setWordWrap(True)
        main.addWidget(explanation)

        controls = QHBoxLayout()
        controls.addWidget(QLabel("Corte institucional:"))
        self._date = QDateEdit()
        self._date.setCalendarPopup(True)
        self._date.setDisplayFormat("dd/MM/yyyy")
        self._date.setToolTip(
            "Corte sugerido: último fin de mes calendario. "
            "La disponibilidad del XML y Maestro debe confirmarse con fuentes reales."
        )
        self._date.setDate(QDate(self._cutoff.year, self._cutoff.month, self._cutoff.day))
        self._date.dateChanged.connect(self._on_date_changed)
        controls.addWidget(self._date)
        controls.addWidget(QLabel("Escenario:"))
        self._scenario = QComboBox()
        self._scenario.setObjectName("sugef210ScenarioSelection")
        self._scenario.addItems(("BASE", "IDIOSYNCRATIC", "SYSTEMIC", "COMBINED"))
        self._scenario.currentTextChanged.connect(self._on_scenario_changed)
        controls.addWidget(self._scenario)
        self._button = QPushButton("Inspeccionar XML de inversiones")
        self._button.clicked.connect(self.run_source_audit)
        controls.addWidget(self._button)
        self._contractual_button = QPushButton("Conciliar captaciones y obligaciones XML")
        self._contractual_button.setObjectName("sugef210RunContractualReconciliation")
        self._contractual_button.clicked.connect(self.run_contractual_audit)
        controls.addWidget(self._contractual_button)
        self._export = QPushButton("Exportar conciliación a Excel")
        self._export.setEnabled(False)
        self._export.clicked.connect(self.export_reconciliation)
        controls.addWidget(self._export)
        self._plan_button = QPushButton("Cargar planificación financiera 2026")
        self._plan_button.clicked.connect(self.select_financial_plan)
        controls.addWidget(self._plan_button)
        self._opening_button = QPushButton("Validar caja inicial XML + Tesorería/ICL")
        self._opening_button.setObjectName("sugef210LoadApprovedOpening")
        self._opening_button.clicked.connect(self.select_governed_opening)
        controls.addWidget(self._opening_button)
        controls.addStretch(1)
        main.addLayout(controls)

        kpi_bar = QHBoxLayout()
        self._projected_kpis: dict[str, QLabel] = {}
        for currency in ("CRC", "USD"):
            frame = QFrame()
            frame.setObjectName(f"sugef210LiquidityKpi{currency}")
            frame.setStyleSheet(
                "QFrame { background-color:#07569A; border-radius:8px; } "
                "QLabel { color:white; background:transparent; }"
            )
            kpi_layout = QVBoxLayout(frame)
            title = QLabel(f"Saldo proyectado parcial a 365 días · {currency}")
            title.setStyleSheet("font-weight:600; font-size:10pt;")
            amount = QLabel("N/D")
            amount.setStyleSheet("font-size:18pt; font-weight:700;")
            amount.setObjectName(f"sugef210BalanceKpi{currency}")
            kpi_layout.addWidget(title)
            kpi_layout.addWidget(amount)
            kpi_bar.addWidget(frame, 1)
            self._projected_kpis[currency] = amount
        main.addLayout(kpi_bar)

        self._status = QLabel(
            "Proyección integral no certificada. Seleccione corte y audite XML real."
        )
        self._status.setWordWrap(True)
        self._status.setStyleSheet(
            "padding:8px; background:#EAF3FA; color:#174E78;" "border-radius:6px; font-weight:600;"
        )
        main.addWidget(self._status)
        self._plan_status = QLabel(
            "Plan financiero: archivo no cargado. Sus supuestos NO se activan automáticamente."
        )
        self._plan_status.setWordWrap(True)
        self._plan_status.setStyleSheet(
            "padding:7px; background:#EDF4FA; color:#064F87; border-radius:6px;"
        )
        main.addWidget(self._plan_status)
        self._opening_status = QLabel(
            "Caja inicial: sin conciliación aprobada de inversiones a la vista. "
            "Ningún saldo XML se añade automáticamente."
        )
        self._opening_status.setWordWrap(True)
        self._opening_status.setStyleSheet(
            "padding:7px; background:#EDF4FA; color:#064F87; border-radius:6px;"
        )
        main.addWidget(self._opening_status)

        splitter = QSplitter(Qt.Orientation.Vertical)
        self._windows = self._table(self._WINDOW_HEADERS, "sugef210LiquidityWindows")
        self._accounts = self._table(self._SOURCE_HEADERS, "sugef210InvestmentXMLAccounts")
        splitter.addWidget(self._windows)
        splitter.addWidget(self._accounts)
        self._records = self._table(self._DETAIL_HEADERS, "sugef210XMLRecordEvidence")
        splitter.addWidget(self._records)
        self._planning_tabs = QTabWidget()
        self._source_readiness = self._table(
            self._SOURCE_READINESS_HEADERS, "sugef210RealSourceReadiness"
        )
        self._planning_tabs.addTab(self._source_readiness, "Estado de fuentes")
        self._contractual_reconciliation = self._table(
            ("Control contractual", "Estado", "Evidencia sin efectivo"),
            "sugef210ContractualReconciliationDetail",
        )
        self._planning_tabs.addTab(self._contractual_reconciliation, "Conciliación contractual")
        self._plan_months = self._table(self._PLAN_HEADERS, "sugef210FinancialPlanMonthlyEvidence")
        self._plan_parameters = self._table(
            self._PLAN_PARAMETER_HEADERS, "sugef210FinancialPlanParameters"
        )
        self._planning_tabs.addTab(self._plan_months, "Plan financiero · calendario")
        self._planning_tabs.addTab(self._plan_parameters, "Plan financiero · supuestos")
        self._xml_perimeters = self._table(
            self._XML_PERIMETER_HEADERS, "sugef210ContractualXMLPerimeters"
        )
        self._planning_tabs.addTab(self._xml_perimeters, "Cartera · captaciones · obligaciones")
        self._governed_flow_details = self._table(
            self._GOVERNED_FLOW_DETAIL_HEADERS, "sugef210GovernedFlowBreakdown"
        )
        self._planning_tabs.addTab(self._governed_flow_details, "Detalle de flujo · fuentes")
        self._scenario_comparison = self._table(
            (
                "Escenario",
                "Banda",
                "Moneda",
                "Δ flujo neto vs BASE",
                "Δ caja proyectada vs BASE",
                "Δ necesidad de financiamiento vs BASE",
                "Estado",
            ),
            "sugef210FourScenarioComparison",
        )
        self._planning_tabs.addTab(self._scenario_comparison, "Comparación 4 escenarios")
        self._investment_reconciliation = self._table(
            self._INVESTMENT_RECONCILIATION_HEADERS,
            "sugef210InvestmentOpeningPrincipalCouponReconciliation",
        )
        self._planning_tabs.addTab(self._investment_reconciliation, "Conciliación inversiones")
        self._event_audit_table = self._table(
            self._EVENT_AUDIT_HEADERS, "sugef210IndividualEventAudit"
        )
        self._planning_tabs.addTab(self._event_audit_table, "Trazabilidad por operación")
        self._demand_deposit_reconciliation = self._table(
            (
                "Cuenta contable XML210",
                "Tipo original",
                "Código moneda original",
                "Registros conciliados",
                "Saldo nativo conciliado",
                "Estado",
                "SHA256 fuente",
            ),
            "sugef210DemandDepositOriginalCurrencyReconciliation",
        )
        self._planning_tabs.addTab(
            self._demand_deposit_reconciliation,
            "Conciliación captaciones vista",
        )
        splitter.addWidget(self._planning_tabs)
        splitter.setSizes([235, 150, 210, 220])
        main.addWidget(splitter, 1)
        self._fill_windows()
        self._fill_operational_gates()

        note = QLabel(
            "Importante: Crediticio_InversionesActivas.xml se utiliza para detectar "
            "inversiones en cuentas corrientes no incluidas en el Maestro. "
            "Sus saldos no son ingresos futuros ni caja disponible hasta "
            "confirmar cuenta contable, conciliación, restricciones y retiro efectivo. "
            "CRC y USD permanecen separados. Los eventos fuera de balance requieren fuente "
            "independiente y los supuestos deben estar documentados y aprobados."
        )
        note.setWordWrap(True)
        main.addWidget(note)

    @staticmethod
    def _table(headers: tuple[str, ...], name: str) -> QTableWidget:
        table = QTableWidget(0, len(headers))
        table.setObjectName(name)
        table.setHorizontalHeaderLabels(headers)
        table.setAlternatingRowColors(True)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        table.setWordWrap(False)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        return table

    @staticmethod
    def _append(table: QTableWidget, values: tuple[str, ...]) -> None:
        row = table.rowCount()
        table.insertRow(row)
        for col, value in enumerate(values):
            item = QTableWidgetItem(value)
            item.setToolTip(value)
            if col in (2, 3, 5, 6):
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            table.setItem(row, col, item)

    def _fill_windows(self) -> None:
        self._windows.setRowCount(0)
        for window in sugef210_windows(self._cutoff):
            self._append(
                self._windows,
                (
                    window.label,
                    window.start.strftime("%d/%m/%Y"),
                    window.end.strftime("%d/%m/%Y"),
                    window.frequency,
                    "N/D",
                    "N/D",
                    "Fuente incompleta",
                    "N/D",
                    "N/D",
                    "N/D",
                    "N/D",
                ),
            )

    def _fill_master_maturities(self, report: dict[str, object]) -> None:
        """Show dated original principal evidence, never mislabel it net cash."""
        maturity = report.get("master_maturity_evidence")
        if not isinstance(maturity, dict):
            return
        if maturity.get("status") != "MASTER_PRINCIPAL_ONLY_PARTIAL_LIQUIDITY_EVIDENCE":
            return
        source_windows = maturity.get("maturity_windows")
        readable_windows = source_windows if isinstance(source_windows, (list, tuple)) else ()
        amounts = {
            (str(row.get("key")), str(row.get("currency"))): str(
                row.get("principal_maturing", "N/D")
            )
            for row in readable_windows
            if isinstance(row, dict)
        }
        for index, window in enumerate(sugef210_windows(self._cutoff)):
            for column, currency in ((7, "CRC"), (8, "USD")):
                value = amounts.get((window.key, currency), "N/D")
                item = QTableWidgetItem(value)
                item.setToolTip(
                    "Principal con fecha de vencimiento verificable en Maestro. "
                    "NO representa una entrada líquida certificada ni incluye cupones."
                )
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self._windows.setItem(index, column, item)

    def _fill_contractual_xml_perimeters(self, report: dict[str, object]) -> None:
        self._xml_perimeters.setRowCount(0)
        raw = report.get("contractual_source_inventory")
        rows = raw if isinstance(raw, (tuple, list)) else ()
        for entry in rows:
            if not isinstance(entry, dict):
                continue
            native = entry.get("source_currency_balances_not_cashflows")
            balances = native if isinstance(native, dict) else {}
            dates = entry.get("dates")
            date_rows = dates if isinstance(dates, (tuple, list)) else ()
            date_text = "; ".join(
                f"{item.get('field')}: {item.get('valid')}/{item.get('missing')}/"
                f"{item.get('invalid')}"
                for item in date_rows
                if isinstance(item, dict)
            )
            self._append(
                self._xml_perimeters,
                (
                    str(entry.get("source_key") or "N/D"),
                    str(entry.get("file") or "NO DISPONIBLE"),
                    str(entry.get("records") if entry.get("file") else "N/D"),
                    ", ".join(f"{key}: {value}" for key, value in balances.items()) or "N/D",
                    date_text or "N/D",
                    str(entry.get("status") or "SIN CERTIFICACIÓN"),
                ),
            )

    def _fill_operational_gates(
        self,
        *,
        xml_status: str = "NOT_INSPECTED",
        master_status: str = "NOT_INSPECTED",
        approved: bool = False,
        projected: bool = False,
        approval_failed: bool = False,
        contractual_inventory: object = None,
    ) -> None:
        """Represent source readiness, never comprehensive liquidity certification."""
        gates = build_operational_gates(
            xml_status=xml_status,
            master_status=master_status,
            approvals_provided=approved,
            projection_executed=projected,
            approval_failed=approval_failed,
            contractual_inventory=contractual_inventory,
        )
        self._source_readiness.setRowCount(0)
        for gate in gates:
            self._append(
                self._source_readiness,
                (gate["check"], gate["state"], gate["next_action"]),
            )

    def _on_scenario_changed(self, _scenario: str) -> None:
        self._invalidate_governed_projection()
        report = self._report or {}
        self._fill_operational_gates(
            xml_status=str(report.get("status") or "NOT_INSPECTED"),
            master_status=str(report.get("master_review_status") or "NOT_INSPECTED"),
            contractual_inventory=report.get("contractual_source_inventory"),
        )
        self._opening_status.setText("Escenario cambiado: validar de nuevo caja inicial.")
        self._status.setText(
            "Escenario cambiado: sus flujos deben estar aprobados por separado. "
            "Proyección anterior invalidada; los datos pendientes permanecen N/D."
        )

    def _invalidate_governed_projection(self) -> None:
        self._governed_projection = None
        self._audited_events = ()
        self._event_audit_table.setRowCount(0)
        self._governed_flow_details.setRowCount(0)
        self._demand_deposit_reconciliation.setRowCount(0)
        self._investment_reconciliation.setRowCount(0)
        self._fill_windows()
        for amount in self._projected_kpis.values():
            amount.setText("N/D")

    def present_governed_projection(self, result: LiquidityProjectionResult) -> None:
        """Present only a caller-supplied governed run, never create assumptions."""
        if result.cutoff != self._cutoff or result.scenario != self._scenario.currentText():
            raise ValueError("Projection cutoff/scenario does not match active Liquidez view")
        if result.windows != sugef210_windows(self._cutoff):
            raise ValueError("Projection horizons differ from fixed SUGEF210 365-day bands")
        expected = {
            (window.key, currency) for window in result.windows for currency in ("CRC", "USD")
        }
        observed = {(row.window.key, row.currency) for row in result.rows}
        if observed != expected or len(result.rows) != len(expected):
            raise ValueError("Projection currencies and horizon rows do not reconcile")
        self._invalidate_governed_projection()
        self._governed_projection = result
        self._fill_governed_flow_details(result)
        if self._report is not None:
            self._fill_master_maturities(self._report)
        self._export.setEnabled(True)
        has_events = {
            currency: any(row.event_count > 0 for row in result.rows if row.currency == currency)
            for currency in ("CRC", "USD")
        }
        latest: dict[str, str] = {"CRC": "N/D", "USD": "N/D"}
        for row in result.rows:
            if row.window.key not in {window.key for window in result.windows}:
                raise ValueError("Unrecognized projected liquidity band")
            index = next(
                (
                    number
                    for number, window in enumerate(result.windows)
                    if window.key == row.window.key
                ),
                None,
            )
            if index is None or index >= self._windows.rowCount():
                raise ValueError("Unrecognized projected liquidity band position")
            net_column = 4 if row.currency == "CRC" else 5
            balance_column = 9 if row.currency == "CRC" else 10
            net = (
                str(row.net_flow)
                if (row.event_count > 0 or result.certified_comprehensive)
                and has_events[row.currency]
                else "N/D"
            )
            balance = str(row.projected_liquidity) if row.projected_liquidity is not None else "N/D"
            net_item = QTableWidgetItem(net)
            net_item.setToolTip(
                "Suma de eventos documentados en esta banda; las fuentes faltantes "
                "no se interpretan como cero. No representa cobertura integral."
            )
            self._windows.setItem(index, net_column, net_item)
            balance_item = QTableWidgetItem(balance)
            balance_item.setToolTip(
                "Saldo inicial verificado más eventos documentados; "
                "excluye capacidad potencial de financiamiento."
            )
            self._windows.setItem(index, balance_column, balance_item)
            self._windows.setItem(
                index,
                6,
                QTableWidgetItem(
                    "CERTIFICADO" if result.certified_comprehensive else "PARCIAL - N/D"
                ),
            )
            if index == self._windows.rowCount() - 1:
                latest[row.currency] = balance
        for currency, label in self._projected_kpis.items():
            label.setText(latest[currency])
        self._status.setText(
            f"Escenario {result.scenario}, corte {result.cutoff.isoformat()}: "
            f"{result.source_flow_count} eventos documentados; "
            f"{result.excluded_out_of_horizon} fuera de 365 días. "
            f"Estado: {result.status}. N/D indica datos aún no certificados."
        )

    def _fill_governed_flow_details(self, result: LiquidityProjectionResult) -> None:
        self._governed_flow_details.setRowCount(0)
        for row in result.rows:
            available = row.event_count > 0 or result.certified_comprehensive
            source_amounts = (
                row.asset_inflows,
                row.asset_outflows,
                row.liability_inflows,
                row.liability_outflows,
                row.off_balance_inflows,
                row.off_balance_outflows,
                row.other_inflows,
                row.other_outflows,
            )
            values = (
                row.window.label,
                row.currency,
                *(
                    str(amount) if amount != 0 or result.certified_comprehensive else "N/D"
                    for amount in source_amounts
                ),
                str(row.net_flow) if available else "N/D",
                str(row.accumulated_flow) if result.source_flow_count else "N/D",
                (str(row.projected_liquidity) if row.projected_liquidity is not None else "N/D"),
                str(row.event_count),
                "INTEGRAL" if result.certified_comprehensive else "PARCIAL",
            )
            self._append(self._governed_flow_details, values)

    def _on_date_changed(self, value: QDate) -> None:
        selected_date = value.toPython()
        if not isinstance(selected_date, date):
            raise TypeError("La fecha de corte seleccionada no es válida")
        self._cutoff = selected_date
        self._scenario_sessions.clear()
        self._scenario_comparison.setRowCount(0)
        self._invalidate_governed_projection()
        self._accounts.setRowCount(0)
        self._records.setRowCount(0)
        self._xml_perimeters.setRowCount(0)
        self._report = None
        self._export.setEnabled(False)
        self._fill_operational_gates()
        self._contractual_reconciliation.setRowCount(0)
        self._plan_audit = None
        self._plan_months.setRowCount(0)
        self._plan_parameters.setRowCount(0)
        self._plan_status.setText(
            "Plan del corte anterior invalidado. Volver a cargar para este corte."
        )
        self._opening_status.setText(
            "Corte cambiado: autorizaciones y conciliaciones de caja anteriores inválidas."
        )
        self._status.setText(
            "Corte cambiado: auditoría pendiente. No se reutilizan cifras del corte anterior."
        )
        self._fill_windows()

    def run_contractual_audit(self) -> None:
        """Read real CONFÍA liabilities without blocking the PySide6 event loop."""
        if (
            self._contractual_worker is not None
            or self._worker is not None
            or self._opening_worker is not None
        ):
            return
        requested_cutoff = self._cutoff
        worker = _SUGEF210ContractualWorker(requested_cutoff)
        self._contractual_worker = worker
        self._contractual_button.setEnabled(False)
        self._button.setEnabled(False)
        self._opening_button.setEnabled(False)
        self._contractual_reconciliation.setRowCount(0)
        self._status.setText(
            "Conciliando claves, principal y fechas de los XML reales en segundo plano. "
            "No se reconocen flujos contractuales sin autorización."
        )
        worker.completed.connect(
            lambda report: self._receive_contractual_audit(requested_cutoff, report)
        )
        worker.failed.connect(
            lambda category: self._contractual_audit_error(requested_cutoff, category)
        )
        worker.finished.connect(self._release_contractual_worker)
        worker.start()

    def _receive_contractual_audit(self, cutoff: date, report: object) -> None:
        if cutoff != self._cutoff:
            return
        try:
            if not isinstance(report, dict):
                raise TypeError("Contractual XML audit must be a structured result")
            rows = summarize_contractual_reconciliation(cutoff=cutoff, evidence=report)
            self._contractual_reconciliation.setRowCount(0)
            for row in rows:
                self._append(self._contractual_reconciliation, row)
            flagged = any(row[1] in {"REVISAR", "NO DISPONIBLE"} for row in rows)
            self._status.setText(
                "Conciliación contractual leída de XML reales. "
                + (
                    "Hay excepciones y posibles fechas operativas por confirmar. "
                    if flagged
                    else "No se detectaron excepciones en los conteos examinados. "
                )
                + "La proyección integral permanece sin certificar."
            )
            self._planning_tabs.setCurrentWidget(self._contractual_reconciliation)
        except Exception as exc:
            self._contractual_audit_error(cutoff, type(exc).__name__)

    def _contractual_audit_error(self, cutoff: date, category: str) -> None:
        if cutoff == self._cutoff:
            self._contractual_reconciliation.setRowCount(0)
            self._status.setText(
                f"Auditoría contractual bloqueada ({category}); sin datos sustitutivos."
            )

    def _release_contractual_worker(self) -> None:
        worker = self._contractual_worker
        if worker is not None and worker.isFinished():
            self._contractual_worker = None
            worker.deleteLater()
            self._contractual_button.setEnabled(True)
            if self._worker is None:
                self._button.setEnabled(True)
            if self._opening_worker is None:
                self._opening_button.setEnabled(True)

    def run_source_audit(self) -> None:
        if (
            self._worker is not None
            or self._opening_worker is not None
            or self._contractual_worker is not None
        ):
            return
        self._invalidate_governed_projection()
        self._opening_status.setText("Auditoría actualizada: volver a validar caja inicial.")
        self._fill_operational_gates()
        requested_cutoff = self._cutoff
        worker = _SUGEF210SourceWorker(requested_cutoff)
        self._worker = worker
        self._button.setEnabled(False)
        self._opening_button.setEnabled(False)
        self._contractual_button.setEnabled(False)
        self._status.setText("Consultando XML institucional; no se calculan flujos sin evidencia…")
        worker.completed.connect(
            lambda report: self._consume_source_audit(requested_cutoff, report)
        )
        worker.failed.connect(lambda category: self._error(requested_cutoff, category))
        worker.finished.connect(self._release_worker)
        worker.start()

    def _consume_source_audit(self, cutoff: date, report: object) -> None:
        """Keep malformed institutional inventory results out of the Qt event loop."""
        if cutoff != self._cutoff:
            return
        try:
            if not isinstance(report, dict):
                raise TypeError("Institutional source inventory must be an object")
            self._receive(cutoff, report)
        except Exception as exc:
            # Display the error category only, never the private source value.
            self._error(cutoff, type(exc).__name__)

    def _receive(self, cutoff: date, report: dict[str, object]) -> None:
        if cutoff != self._cutoff:
            return
        self._accounts.setRowCount(0)
        self._records.setRowCount(0)
        self._report = None
        self._export.setEnabled(False)
        status = str(report.get("status", "BLOCKED"))
        self._fill_operational_gates(
            xml_status=status,
            master_status=str(report.get("master_review_status") or "NOT_INSPECTED"),
            contractual_inventory=report.get("contractual_source_inventory"),
        )
        if status != "ACCOUNT_LEVEL_EVIDENCE_ONLY_NO_AVAILABLE_CASH":
            self._status.setText(f"Fuente de inversiones no utilizable para inspección: {status}")
            return
        self._fill_master_maturities(report)
        self._fill_contractual_xml_perimeters(report)
        account_source = report.get("accounts")
        account_rows = account_source if isinstance(account_source, (list, tuple)) else ()
        for group in account_rows:
            if not isinstance(group, dict):
                continue
            self._append(
                self._accounts,
                (
                    str(group["account_code"]),
                    str(group["currency_source_code"]),
                    str(group["record_count"]),
                    str(group["source_principal"]),
                    str(group["classification"]),
                    str(group["master_isin_matches"]),
                    str(group["unverified_master_overlap"]),
                ),
            )
        detail_source = report.get("record_evidence")
        detail_rows = detail_source if isinstance(detail_source, (list, tuple)) else ()
        for record in detail_rows:
            if not isinstance(record, dict):
                continue
            self._append(
                self._records,
                (
                    str(record.get("source_reference") or ""),
                    str(record.get("raw_record_id") or ""),
                    str(record.get("account_code") or ""),
                    str(record.get("currency_source_code") or ""),
                    str(record.get("source_principal") or "N/D"),
                    str(record.get("instrument_id") or ""),
                    str(record.get("isin") or ""),
                    str(record.get("series") or ""),
                    str(record.get("master_reconciliation") or ""),
                    str(record.get("classification") or ""),
                ),
            )
        self._report = report
        self._export.setEnabled(True)
        master_status = str(report.get("master_review_status") or "N/D")
        source_step = (
            "Maestro exacto del corte identificado; para ejecutar la proyección "
            "seleccione autorización XML y conciliación Tesorería/ICL vigentes."
            if master_status == "MASTER_EXACT_CUTOFF_READ"
            else "Maestro exacto no validado: revise fecha de corte y ruta del portafolio "
            "antes de cargar saldos o vencimientos."
        )
        self._status.setText(
            f"XML inversiones: {report.get('xml_records_read', 0)} registros; "
            f"{self._accounts.rowCount()} cuentas/monedas inspeccionadas; "
            f"{self._records.rowCount()} registros trazables. "
            f"Maestro: {master_status}. {source_step} "
            "Disponibilidad a la vista: NO CERTIFICADA. "
            "Flujo integral: pendiente de fuentes y supuestos institucionales."
        )

    def _error(self, cutoff: date, category: str) -> None:
        if cutoff == self._cutoff:
            self._invalidate_governed_projection()
            self._accounts.setRowCount(0)
            self._records.setRowCount(0)
            self._xml_perimeters.setRowCount(0)
            self._report = None
            self._export.setEnabled(False)
            self._fill_operational_gates(
                xml_status="BLOCKED_SOURCE_READ", master_status="NOT_INSPECTED"
            )
            self._status.setText(
                f"BLOCKED_SOURCE_READ ({category}). Sin sustituir información ni mostrar flujos."
            )

    def select_governed_opening(self) -> None:
        """Run real configured Master/XML and two separate institutional approvals."""
        if (
            self._opening_worker is not None
            or self._worker is not None
            or self._contractual_worker is not None
        ):
            return
        approval, _kind = QFileDialog.getOpenFileName(
            self, "Autorización de inversiones a la vista SUGEF210", "", "JSON (*.json)"
        )
        if not approval:
            return
        bank, _kind = QFileDialog.getOpenFileName(
            self, "Conciliación de caja Tesorería/ICL SUGEF210", "", "JSON (*.json)"
        )
        if not bank:
            return
        cutoff = self._cutoff
        scenario = self._scenario.currentText()
        self._scenario_sessions.pop(scenario, None)
        self._scenario_comparison.setRowCount(0)
        self._invalidate_governed_projection()
        self._opening_button.setEnabled(False)
        self._button.setEnabled(False)
        self._contractual_button.setEnabled(False)
        self._opening_status.setText(
            "Validando XML real, Maestro, cuenta a la vista y respaldo bancario. "
            "No se aceptan saldos sin trazabilidad."
        )
        worker = _SUGEF210OpeningWorker(cutoff, scenario, approval, bank)
        self._opening_worker = worker
        worker.completed.connect(
            lambda session: self._consume_governed_opening(cutoff, scenario, session)
        )
        worker.failed.connect(
            lambda category: self._governed_opening_error(cutoff, scenario, category)
        )
        worker.finished.connect(self._release_opening_worker)
        worker.start()

    def _investment_reconciliation_rows(
        self, session: ConfiguredOpeningSession
    ) -> tuple[tuple[str, ...], ...]:
        """Validate original-currency source evidence before updating the UI."""
        if session.cutoff != self._cutoff or session.scenario != self._scenario.currentText():
            raise ValueError("Investment reconciliation cutoff/scenario mismatch")
        by_source = {
            (row.source_name, row.currency): row
            for row in session.assembly.source_reconciliations
            if row.source_name in {"MASTER_RECOGNIZED_PRINCIPAL", "APPROVED_MASTER_COUPONS"}
        }
        if len(by_source) != sum(
            1
            for row in session.assembly.source_reconciliations
            if row.source_name in {"MASTER_RECOGNIZED_PRINCIPAL", "APPROVED_MASTER_COUPONS"}
        ):
            raise ValueError("Duplicated investment source/currency reconciliation")
        source_counts = {
            source: sum(
                row.received_events
                for (name, _currency), row in by_source.items()
                if name == source
            )
            for source in ("MASTER_RECOGNIZED_PRINCIPAL", "APPROVED_MASTER_COUPONS")
        }
        if (
            source_counts["MASTER_RECOGNIZED_PRINCIPAL"] != session.master_principal_events_approved
            or source_counts["APPROVED_MASTER_COUPONS"] != session.master_coupon_events_approved
        ):
            raise ValueError("Investment event counts differ from source reconciliations")
        sight_amounts = dict(session.sight_added_by_currency)
        if any(
            currency not in {"CRC", "USD"} or not amount.is_finite() or amount < 0
            for currency, amount in sight_amounts.items()
        ) or len(sight_amounts) != len(session.sight_added_by_currency):
            raise ValueError("Original-currency sight investment additions invalid")
        rows: list[tuple[str, ...]] = []
        for currency in ("CRC", "USD"):
            sight = sight_amounts.get(currency)
            rows.append(
                (
                    "Inversiones a la vista XML",
                    "APERTURA · no flujo futuro",
                    currency,
                    "N/D",
                    "N/D",
                    str(sight) if sight is not None else "N/D",
                    "APERTURA VERIFICADA · PARCIAL" if sight is not None else "N/D · SIN DESGLOSE",
                    session.sight_source_sha256 if sight is not None else "N/D",
                ),
            )
            for name, label, digest in (
                (
                    "MASTER_RECOGNIZED_PRINCIPAL",
                    "Principal Maestro · futuro",
                    session.master_source_sha256,
                ),
                (
                    "APPROVED_MASTER_COUPONS",
                    "Cupones Custodia · futuro",
                    session.master_coupon_custodian_sha256,
                ),
            ):
                reconciled = by_source.get((name, currency))
                rows.append(
                    (
                        name,
                        label,
                        currency,
                        str(reconciled.events_inside_horizon) if reconciled else "N/D",
                        str(reconciled.events_after_horizon) if reconciled else "N/D",
                        (
                            str(reconciled.principal_or_flow_total_inside_horizon)
                            if reconciled
                            else "N/D"
                        ),
                        "PARCIAL · " + reconciled.status if reconciled else "N/D · NO INCORPORADO",
                        digest if reconciled and digest else "N/D",
                    ),
                )
        return tuple(rows)

    @staticmethod
    def _event_audit_row(event: LiquidityAuditedCashEvent) -> tuple[str, ...]:
        return (
            event.source_group,
            event.reference,
            event.source_reference,
            event.payment_date.isoformat(),
            event.currency,
            str(event.native_amount),
            event.area.value,
            event.side.value,
            event.evidence.value,
            event.assumption_id or "N/D",
            event.band_key,
            event.horizon_status,
            event.source_status,
        )

    def _validated_event_audit(
        self, session: ConfiguredOpeningSession
    ) -> tuple[LiquidityAuditedCashEvent, ...]:
        """Do not render tampered/unreconciled cash-event lineage."""
        events = session.assembly.audited_events
        projection = session.assembly.projection
        if len(events) != projection.source_flow_count:
            raise ValueError("Event lineage count differs from approved projection")
        if (
            sum(e.horizon_status == "AFTER_365D" for e in events)
            != projection.excluded_out_of_horizon
        ):
            raise ValueError("Event lineage beyond 365 days differs from projection")
        if len({e.reference for e in events}) != len(events):
            raise ValueError("Duplicated institutional cash-event lineage")
        # O(events + 34 bands): retain a second fail-closed verification
        # after the worker returns, without repeatedly scanning the full ledger.
        valid_windows = {window.key: window for window in projection.windows}
        audit_counts: dict[tuple[str, str], int] = {}
        audit_amounts: dict[tuple[str, str, LiquidityBalanceArea, LiquidityFlowSide], Decimal] = {}
        for event in events:
            if (
                not event.native_amount.is_finite()
                or event.native_amount <= 0
                or event.currency not in {"CRC", "USD"}
            ):
                raise ValueError("Invalid source amount or currency in event audit")
            if event.horizon_status == "AFTER_365D":
                if (
                    event.band_key != "OUTSIDE_365D"
                    or event.payment_date <= projection.windows[-1].end
                ):
                    raise ValueError("Cash-event outside-horizon classification inconsistent")
            elif event.horizon_status == "IN_365D":
                window = valid_windows.get(event.band_key)
                if window is None or not window.start <= event.payment_date <= window.end:
                    raise ValueError("Cash-event band classification inconsistent")
                count_key = (event.currency, event.band_key)
                audit_counts[count_key] = audit_counts.get(count_key, 0) + 1
                key = (event.currency, event.band_key, event.area, event.side)
                audit_amounts[key] = audit_amounts.get(key, Decimal(0)) + event.native_amount
            else:
                raise ValueError("Unknown cash-event horizon classification")
        for row in projection.rows:
            if audit_counts.get((row.currency, row.window.key), 0) != row.event_count:
                raise ValueError("Event lineage does not reconcile to 34 projected bands")
            categories = (
                (LiquidityBalanceArea.ASSET, LiquidityFlowSide.INFLOW, row.asset_inflows),
                (LiquidityBalanceArea.ASSET, LiquidityFlowSide.OUTFLOW, row.asset_outflows),
                (LiquidityBalanceArea.LIABILITY, LiquidityFlowSide.INFLOW, row.liability_inflows),
                (LiquidityBalanceArea.LIABILITY, LiquidityFlowSide.OUTFLOW, row.liability_outflows),
                (
                    LiquidityBalanceArea.OFF_BALANCE,
                    LiquidityFlowSide.INFLOW,
                    row.off_balance_inflows,
                ),
                (
                    LiquidityBalanceArea.OFF_BALANCE,
                    LiquidityFlowSide.OUTFLOW,
                    row.off_balance_outflows,
                ),
                (LiquidityBalanceArea.BUSINESS_EVENT, LiquidityFlowSide.INFLOW, row.other_inflows),
                (
                    LiquidityBalanceArea.BUSINESS_EVENT,
                    LiquidityFlowSide.OUTFLOW,
                    row.other_outflows,
                ),
            )
            for area, direction, expected in categories:
                key = (row.currency, row.window.key, area, direction)
                if audit_amounts.get(key, Decimal(0)) != expected:
                    raise ValueError("Event lineage native amounts differ from projection")
        return events

    def _consume_governed_opening(
        self, cutoff: date, scenario: str, session: ConfiguredOpeningSession
    ) -> None:
        """Reject mismatched event evidence without raising through a Qt signal."""
        if cutoff != self._cutoff or scenario != self._scenario.currentText():
            return
        try:
            self._receive_governed_opening(cutoff, scenario, session)
        except Exception as exc:
            # Never print paths or amounts from confidential institution files.
            self._governed_opening_error(cutoff, scenario, type(exc).__name__)

    def _receive_governed_opening(
        self, cutoff: date, scenario: str, session: ConfiguredOpeningSession
    ) -> None:
        if cutoff != self._cutoff or scenario != self._scenario.currentText():
            return
        if session.cutoff != cutoff or session.scenario != scenario:
            raise ValueError("Loaded opening cutoff or scenario does not match the request")
        # Fail closed on source-level mismatches before changing the visible forecast.
        investment_rows = self._investment_reconciliation_rows(session)
        audited_events = self._validated_event_audit(session)
        self.present_governed_projection(session.assembly.projection)
        self._audited_events = audited_events
        self._event_audit_table.setUpdatesEnabled(False)
        try:
            for event in audited_events[: self._EVENT_AUDIT_UI_LIMIT]:
                self._append(self._event_audit_table, self._event_audit_row(event))
        finally:
            self._event_audit_table.setUpdatesEnabled(True)
        self._status.setText(
            self._status.text()
            + f" Trazabilidad: {min(len(audited_events), self._EVENT_AUDIT_UI_LIMIT)}"
            + f" de {len(audited_events)} operaciones en pantalla; Excel incluye todas."
        )
        self._investment_reconciliation.setRowCount(0)
        for values in investment_rows:
            self._append(self._investment_reconciliation, values)
        self._demand_deposit_reconciliation.setRowCount(0)
        for (
            account,
            original_type,
            native_code,
            count,
            original_balance,
        ) in session.demand_deposit_native_reconciliations:
            self._append(
                self._demand_deposit_reconciliation,
                (
                    account,
                    original_type,
                    native_code,
                    str(count),
                    str(original_balance),
                    "ORIGEN CONCILIADO · COBERTURA PARCIAL",
                    session.demand_deposit_source_sha256 or "N/D",
                ),
            )
        self._fill_operational_gates(
            xml_status="ACCOUNT_LEVEL_EVIDENCE_ONLY_NO_AVAILABLE_CASH",
            master_status="MASTER_EXACT_CUTOFF_READ",
            approved=True,
            projected=True,
            contractual_inventory=(self._report or {}).get("contractual_source_inventory"),
        )
        for coverage in summarize_approved_contractual_source_coverage(
            session.xml_contractual_source_coverage
        ):
            source_key = str(coverage["source_key"])
            approved_count = coverage["approved_contracts"]
            original_count = coverage["original_xml_contracts"]
            if approved_count is None:
                gate_status = "NO_APROBADO"
                detail = (
                    "Sin cronograma contractual aprobado para esta fuente XML. "
                    "Cobertura y pagos no determinados; no equivale a cero."
                )
            else:
                gate_status = (
                    "REVIEW_REQUIRED" if approved_count != original_count else "REFERENCE_ONLY"
                )
                detail = (
                    f"{approved_count} contratos aprobados de {original_count} "
                    "registros XML originales. Los no seleccionados requieren "
                    "clasificación; cobertura completa no equivale a certificación."
                )
            self._append(
                self._source_readiness,
                (f"XML_CRONOGRAMA_{source_key}", gate_status, detail),
            )
        self._scenario_sessions[scenario] = session
        self._refresh_four_scenario_comparison()
        self._opening_status.setText(
            f"Caja inicial documentada: {session.sight_records_approved} inversiones "
            f"a la vista revisadas contra XML/Maestro; "
            f"{session.master_principal_events_approved} vencimientos de principal del Maestro "
            "autorizados por instrumento; "
            f"{session.master_coupon_events_approved} cupones documentados por custodia "
            "y aprobados individualmente (sin cálculo de tasas); "
            f"{session.xml_contractual_events_approved} pagos individuales "
            "documentados de cartera, captaciones y obligaciones; "
            f"{session.offbalance_commitments_reviewed} compromisos contingentes revisados, "
            f"{session.contingent_scenario_events_approved} desembolsos potenciales; "
            f"{session.financial_plan_events_approved} eventos del plan financiero "
            "con fecha, moneda y conciliación aprobadas; "
            f"{session.demand_deposit_accounts_reviewed} depósitos a la vista revisados, "
            f"{session.demand_deposit_runoff_events_approved} retiros conductuales "
            f"bajo escenario {session.scenario} aprobado. "
            f"XML SHA256 {session.sight_source_sha256[:12]}…; "
            f"respaldo bancario SHA256 {session.bank_document_sha256[:12]}…. "
            "Proyección PARCIAL: contingencias y cronogramas no seleccionados "
            "permanecen sin cobertura integral. "
            "La herramienta no verifica por sí sola firmas ni cifras de certificados bancarios."
        )

    def _refresh_four_scenario_comparison(self) -> None:
        """Compare loaded approvals only when all four share one source snapshot."""
        self._scenario_comparison.setRowCount(0)
        if len(self._scenario_sessions) != 4:
            return
        try:
            comparison = compare_approved_sugef210_scenarios(self._scenario_sessions.values())
        except ValueError:
            self._append(
                self._scenario_comparison,
                (
                    "BLOQUEADO",
                    "N/D",
                    "N/D",
                    "N/D",
                    "N/D",
                    "N/D",
                    "Cortes, fuentes, caja o pagos contractuales no coinciden entre escenarios",
                ),
            )
            return
        self._scenario_comparison.setUpdatesEnabled(False)
        try:
            for row in comparison.rows:
                self._append(
                    self._scenario_comparison,
                    (
                        row.scenario,
                        row.band_key,
                        row.currency,
                        str(row.net_flow_delta_vs_base),
                        (
                            str(row.projected_liquidity_delta_vs_base)
                            if row.projected_liquidity_delta_vs_base is not None
                            else "N/D"
                        ),
                        (
                            str(row.financing_need_delta_vs_base)
                            if row.financing_need_delta_vs_base is not None
                            else "N/D"
                        ),
                        "PARCIAL · NO CERTIFICADO",
                    ),
                )
        finally:
            self._scenario_comparison.setUpdatesEnabled(True)

    def _governed_opening_error(self, cutoff: date, scenario: str, category: str) -> None:
        if cutoff != self._cutoff or scenario != self._scenario.currentText():
            return
        self._scenario_sessions.pop(scenario, None)
        self._scenario_comparison.setRowCount(0)
        self._invalidate_governed_projection()
        report = self._report or {}
        self._fill_operational_gates(
            xml_status=str(report.get("status") or "NOT_INSPECTED"),
            master_status=str(report.get("master_review_status") or "NOT_INSPECTED"),
            approved=True,
            approval_failed=True,
            contractual_inventory=report.get("contractual_source_inventory"),
        )
        self._opening_status.setText(
            f"BLOCKED_GOVERNED_OPENING ({category}). No se incorporó caja XML."
        )

    def _release_opening_worker(self) -> None:
        worker = self._opening_worker
        if worker is not None and worker.isFinished():
            self._opening_worker = None
            worker.deleteLater()
            self._opening_button.setEnabled(True)
            if self._worker is None:
                self._button.setEnabled(True)
            if self._contractual_worker is None:
                self._contractual_button.setEnabled(True)

    def select_financial_plan(self) -> None:
        """Read local institutional workbook without copying it into the Git repository."""
        if self._plan_worker is not None:
            return
        source_path, _file_filter = QFileDialog.getOpenFileName(
            self,
            "Seleccionar Matriz de Supuestos del Flujo de Efectivo 2026",
            "",
            "Excel (*.xlsx)",
        )
        if not source_path:
            return
        cutoff = self._cutoff
        worker = _SUGEF210FinancialPlanWorker(cutoff, source_path)
        self._plan_worker = worker
        self._plan_button.setEnabled(False)
        self._plan_status.setText(
            "Leyendo la matriz institucional en modo seguro. "
            "No se registrarán importes como flujos contractuales."
        )
        worker.completed.connect(lambda evidence: self._consume_financial_plan(cutoff, evidence))
        worker.failed.connect(lambda reason: self._financial_plan_error(cutoff, reason))
        worker.finished.connect(self._release_plan_worker)
        worker.start()

    def _consume_financial_plan(self, cutoff: date, evidence: object) -> None:
        """Reset an invalid workbook payload without an uncaught Qt slot exception."""
        if cutoff != self._cutoff:
            return
        try:
            if not isinstance(evidence, FinancialPlanSourceAudit):
                raise TypeError("Financial plan audit must be validated evidence")
            self._receive_financial_plan(cutoff, evidence)
        except Exception as exc:
            self._financial_plan_error(cutoff, type(exc).__name__)

    def _receive_financial_plan(self, cutoff: date, evidence: FinancialPlanSourceAudit) -> None:
        if cutoff != self._cutoff:
            return
        self._plan_audit = evidence
        self._plan_months.setRowCount(0)
        self._plan_parameters.setRowCount(0)
        for row in evidence.monthly_inputs:
            self._append(
                self._plan_months,
                (
                    f"{row.month:02d}/{row.year}",
                    row.driver,
                    str(row.reported_value),
                    str(row.area or "NO EFECTIVO / PARÁMETRO"),
                    str(row.side or "—"),
                    row.status,
                    row.source_cell,
                ),
            )
        for param in evidence.parameters:
            self._append(
                self._plan_parameters,
                (
                    str(param.plan_id),
                    param.name,
                    param.value if param.value is not None else "N/D",
                    param.unit,
                    param.periodicity,
                    param.cashflow_type,
                    param.plan_reference,
                    param.liquidity_role,
                    param.source_cell,
                ),
            )
        next_months = sum(
            row.period_start > date(cutoff.year, cutoff.month, 1)
            and row.classification == "PLANNED_POTENTIAL_CASHFLOW"
            for row in evidence.monthly_inputs
        )
        self._plan_status.setText(
            f"Plan {evidence.plan_year}: {len(evidence.parameters)} supuestos, "
            f"{next_months} importes mensuales futuros identificados. "
            f"SHA256: {evidence.sha256[:16]}… "
            "Datos de referencia en ¢ millones; NO son CRC efectivo ni USD certificado. "
            "Pendientes: aprobación, fecha, moneda original y no duplicación."
        )

    def _financial_plan_error(self, cutoff: date, reason: str) -> None:
        if cutoff != self._cutoff:
            return
        self._plan_audit = None
        self._plan_months.setRowCount(0)
        self._plan_parameters.setRowCount(0)
        self._plan_status.setText(
            f"BLOCKED_FINANCIAL_PLAN_{reason}: no se incorporaron proyecciones."
        )

    def _release_plan_worker(self) -> None:
        worker = self._plan_worker
        if worker is not None and worker.isFinished():
            self._plan_worker = None
            worker.deleteLater()
            self._plan_button.setEnabled(True)

    def export_reconciliation(self) -> None:
        """Export real audit evidence; never classify a balance as cash."""
        if self._report is None and self._governed_projection is None:
            return
        selected, _filter = QFileDialog.getSaveFileName(
            self,
            "Exportar auditoría XML–Maestro",
            f"SUGEF210_Conciliacion_{self._cutoff.isoformat()}.xlsx",
            "Excel (*.xlsx)",
        )
        if not selected:
            return
        from openpyxl import Workbook

        workbook = Workbook()
        metadata = workbook.active
        assert metadata is not None
        metadata.title = "Control"
        report = self._report or {}
        for key in (
            "cutoff",
            "xml_month",
            "xml_source",
            "entity_id",
            "sha256",
            "master_review_status",
            "master_positions_read",
            "xml_records_read",
            "invalid_source_principal",
            "missing_account_codes",
            "repeated_source_record_ids",
        ):
            metadata.append((key, str(report.get(key, "N/D"))))
        metadata.append(("disponibilidad_certificada", "NO"))
        projection = self._governed_projection
        metadata.append(
            (
                "flujo_integral_certificado",
                "SI" if projection is not None and projection.certified_comprehensive else "NO",
            )
        )
        if projection is not None:
            metadata.append(("escenario", projection.scenario))
            metadata.append(("eventos_documentados", str(projection.source_flow_count)))
            metadata.append(("estado_proyeccion", projection.status))
            metadata.append(("fuentes_faltantes", ", ".join(projection.missing_source_areas)))
            metadata.append(("supuestos_trazables", ", ".join(projection.assumption_ids)))
        metadata.append(("nota", "Identificadores no prueban titularidad ni retiro disponible"))

        def safe(value: object) -> str | Decimal:
            # Preserve monetary values as numeric Excel cells, including negative GAPs.
            # Escape only untrusted strings that Excel could interpret as formulas.
            if isinstance(value, Decimal):
                if not value.is_finite():
                    raise ValueError("Nonfinite liquidity amount cannot be exported")
                return value
            content = "" if value is None else str(value)
            return "'" + content if content.startswith(("=", "+", "-", "@")) else content

        if projection is not None:
            projected = workbook.create_sheet("Flujo proyectado 365d")
            projected.append(
                (
                    "Escenario",
                    "Corte",
                    "Horizonte",
                    "Desde",
                    "Hasta",
                    "Moneda",
                    "Entradas activos",
                    "Salidas activos",
                    "Entradas pasivos",
                    "Salidas pasivos",
                    "Entradas fuera balance",
                    "Salidas fuera balance",
                    "Entradas negocio",
                    "Salidas negocio",
                    "Neto parcial",
                    "GAP acumulado parcial",
                    "Disponibilidades iniciales",
                    "Saldo proyectado parcial",
                    "Necesidad financiación parcial",
                    "Estado",
                )
            )
            for item in projection.rows:
                projected.append(
                    tuple(
                        safe(value)
                        for value in (
                            projection.scenario,
                            projection.cutoff.isoformat(),
                            item.window.label,
                            item.window.start.isoformat(),
                            item.window.end.isoformat(),
                            item.currency,
                            (
                                item.asset_inflows
                                if item.asset_inflows != 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.asset_outflows
                                if item.asset_outflows != 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.liability_inflows
                                if item.liability_inflows != 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.liability_outflows
                                if item.liability_outflows != 0
                                or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.off_balance_inflows
                                if item.off_balance_inflows != 0
                                or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.off_balance_outflows
                                if item.off_balance_outflows != 0
                                or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.other_inflows
                                if item.other_inflows != 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.other_outflows
                                if item.other_outflows != 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            (
                                item.net_flow
                                if item.event_count > 0 or projection.certified_comprehensive
                                else "N/D"
                            ),
                            item.accumulated_flow if projection.source_flow_count else "N/D",
                            item.opening_liquidity if item.opening_liquidity is not None else "N/D",
                            (
                                item.projected_liquidity
                                if item.projected_liquidity is not None
                                else "N/D"
                            ),
                            item.financing_need if item.financing_need is not None else "N/D",
                            projection.status,
                        )
                    )
                )

        comparison_sheet = workbook.create_sheet("Comparación 4 escenarios")
        comparison_sheet.append(
            (
                "Escenario",
                "Corte",
                "Banda",
                "Moneda",
                "Flujo neto",
                "Δ flujo neto vs BASE",
                "Δ GAP acumulado vs BASE",
                "Saldo proyectado",
                "Δ saldo proyectado vs BASE",
                "Necesidad de financiamiento",
                "Δ necesidad vs BASE",
                "Evidencia faltante",
                "Certificación",
            )
        )
        if len(self._scenario_sessions) == 4:
            try:
                compared = compare_approved_sugef210_scenarios(self._scenario_sessions.values())
            except ValueError:
                metadata.append(("comparacion_escenarios", "BLOQUEADA: fuentes no comparables"))
                comparison_sheet.append(
                    ("BLOQUEADA", self._cutoff.isoformat(), "N/D", "N/D", "N/D")
                )
            else:
                metadata.append(
                    ("comparacion_escenarios", "COMPARABLE; NO ES CERTIFICACION REGULATORIA")
                )
                by_scenario = dict(compared.missing_evidence_by_scenario)
                for item in compared.rows:
                    comparison_sheet.append(
                        tuple(
                            safe(value)
                            for value in (
                                item.scenario,
                                compared.cutoff,
                                item.band_key,
                                item.currency,
                                item.net_flow,
                                item.net_flow_delta_vs_base,
                                item.accumulated_flow_delta_vs_base,
                                (
                                    item.projected_liquidity
                                    if item.projected_liquidity is not None
                                    else "N/D"
                                ),
                                (
                                    item.projected_liquidity_delta_vs_base
                                    if item.projected_liquidity_delta_vs_base is not None
                                    else "N/D"
                                ),
                                (item.financing_need if item.financing_need is not None else "N/D"),
                                (
                                    item.financing_need_delta_vs_base
                                    if item.financing_need_delta_vs_base is not None
                                    else "N/D"
                                ),
                                ", ".join(by_scenario.get(item.scenario, ())),
                                "NO CERTIFICADO",
                            )
                        )
                    )
        else:
            metadata.append(
                ("comparacion_escenarios", "INCOMPLETA: requieren cuatro sesiones aprobadas")
            )
            comparison_sheet.append(("INCOMPLETA", self._cutoff.isoformat(), "N/D", "N/D", "N/D"))

        event_sheet = workbook.create_sheet("Trazabilidad eventos")
        event_sheet.append(self._EVENT_AUDIT_HEADERS)
        for audited in self._audited_events:
            audit_values = self._event_audit_row(audited)
            event_sheet.append(
                tuple(
                    audited.native_amount if column == 5 else safe(value)
                    for column, value in enumerate(audit_values)
                )
            )
        if projection is not None:
            metadata.append(("eventos_auditados_linea_a_linea", len(self._audited_events)))
            metadata.append(
                (
                    "eventos_mostrados_en_UI",
                    min(len(self._audited_events), self._EVENT_AUDIT_UI_LIMIT),
                )
            )
            metadata.append(
                ("nota_trazabilidad", "Excel incluye todos los eventos; UI muestra máximo 300")
            )

        investment_sheet = workbook.create_sheet("Conciliacion inversiones")
        investment_sheet.append(self._INVESTMENT_RECONCILIATION_HEADERS)
        for row_index in range(self._investment_reconciliation.rowCount()):
            investment_cells: list[str | int | Decimal] = []
            for column in range(self._investment_reconciliation.columnCount()):
                item = self._investment_reconciliation.item(row_index, column)
                raw = item.text() if item is not None else ""
                if column in (3, 4) and raw != "N/D":
                    investment_cells.append(int(raw))
                elif column == 5 and raw != "N/D":
                    investment_cells.append(Decimal(raw))
                else:
                    investment_cells.append(safe(raw))
            investment_sheet.append(tuple(investment_cells))

        sight_sheet = workbook.create_sheet("Captaciones vista XML210")
        sight_sheet.append(
            (
                "Cuenta contable",
                "Tipo original",
                "Código moneda XML",
                "Registros conciliados",
                "Saldo original (sin convertir)",
                "Estado de cobertura",
                "SHA256 de XML 210",
            )
        )
        for row_index in range(self._demand_deposit_reconciliation.rowCount()):
            native_cells: list[str | Decimal] = []
            for column in range(self._demand_deposit_reconciliation.columnCount()):
                cell = self._demand_deposit_reconciliation.item(row_index, column)
                native_cells.append(safe(cell.text()) if cell is not None else "")
            sight_sheet.append(tuple(native_cells))

        perimeter = workbook.create_sheet("Fuentes contractuales XML")
        perimeter.append(self._XML_PERIMETER_HEADERS)
        for index in range(self._xml_perimeters.rowCount()):
            cells: list[str | Decimal] = []
            for column in range(len(self._XML_PERIMETER_HEADERS)):
                cell = self._xml_perimeters.item(index, column)
                cells.append(safe(cell.text()) if cell is not None else "")
            perimeter.append(tuple(cells))

        accounts = workbook.create_sheet("Cuentas XML")
        account_keys = (
            "account_code",
            "currency_source_code",
            "record_count",
            "source_principal",
            "classification",
            "master_isin_matches",
            "unverified_master_overlap",
        )
        accounts.append(account_keys)
        source_accounts = report.get("accounts")
        account_rows = source_accounts if isinstance(source_accounts, (list, tuple)) else ()
        for row in account_rows:
            if isinstance(row, dict):
                accounts.append(tuple(safe(row.get(field)) for field in account_keys))

        maturity = report.get("master_maturity_evidence")
        if isinstance(maturity, dict):
            schedule = workbook.create_sheet("Vencimientos Maestro")
            keys = ("key", "from", "to", "currency", "principal_maturing", "status")
            schedule.append(keys)
            master_rows = maturity.get("maturity_windows")
            scheduled_rows = master_rows if isinstance(master_rows, (list, tuple)) else ()
            for item in scheduled_rows:
                if isinstance(item, dict):
                    schedule.append(tuple(safe(item.get(key)) for key in keys))

        records = workbook.create_sheet("Detalle conciliación")
        record_keys = (
            "source_reference",
            "raw_record_id",
            "record_ordinal",
            "action",
            "account_code",
            "currency_source_code",
            "source_principal_raw",
            "source_principal",
            "instrument_id",
            "isin",
            "series",
            "classification",
            "master_reconciliation",
            "withdrawable_amount",
        )
        records.append(record_keys)
        detailed_source = report.get("record_evidence")
        detailed_rows = detailed_source if isinstance(detailed_source, (list, tuple)) else ()
        for row in detailed_rows:
            if isinstance(row, dict):
                records.append(tuple(safe(row.get(field)) for field in record_keys))
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        try:
            workbook.save(selected)
        except OSError as exc:
            self._status.setText(
                f"Exportación bloqueada: {type(exc).__name__}. XML y Maestro sin cambios."
            )
            return
        self._status.setText(
            "Conciliación exportada en Excel. Disponibilidad de caja y "
            "diferencias de portafolio pendientes de certificación."
        )

    def _release_worker(self) -> None:
        worker = self._worker
        if worker is not None and worker.isFinished():
            self._worker = None
            worker.deleteLater()
            self._button.setEnabled(True)
            if self._opening_worker is None:
                self._opening_button.setEnabled(True)
            if self._contractual_worker is None:
                self._contractual_button.setEnabled(True)
