"""Excel charts for audited, explicitly partial SUGEF 2-10 cashflow projections.

No N/D value is converted to zero. Charts describe admitted events only,
never assert that all balance-sheet or off-balance-sheet flows are covered.
"""

from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook import Workbook

from aip.domain.liquidity.cashflow.services.sugef210_projection import (
    LiquidityProjectionResult,
)


def add_partial_cashflow_charts(workbook: Workbook, projection: LiquidityProjectionResult) -> None:
    """Plot evidenced native-currency cashflows without inventing missing flows."""
    sheet = workbook.create_sheet("Gráficos parciales")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "SUGEF 2-10 — FLUJOS DOCUMENTADOS; PROYECCIÓN PARCIAL NO CERTIFICADA"
    sheet["A2"] = (
        "N/D no equivale a cero. Cada moneda se grafica por separado; "
        "GAP acumulado incluye solo eventos admitidos."
    )
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    sheet["A1"].fill = PatternFill("solid", fgColor="164B85")
    sheet["A2"].font = Font(italic=True, color="9C4900", size=10)
    sheet.merge_cells("A1:E1")
    sheet.merge_cells("A2:E2")
    sheet["A1"].alignment = Alignment(wrap_text=True)
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 32
    sheet.append(())  # Headers are written explicitly at row 4 below.
    headers = (
        "Banda",
        "CRC neto parcial",
        "CRC GAP acumulado parcial",
        "USD neto parcial",
        "USD GAP acumulado parcial",
    )
    for column, name in enumerate(headers, 1):
        cell = sheet.cell(row=4, column=column, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1760A5")
    sheet.column_dimensions["A"].width = 27
    for letter in "BCDE":
        sheet.column_dimensions[letter].width = 27
    sheet.freeze_panes = "B5"

    indexed = {(row.currency, row.window.key): row for row in projection.rows}
    if len(indexed) != len(projection.rows):
        raise ValueError("Cannot chart duplicate cashflow currency/band entries")
    if len(indexed) != 2 * len(projection.windows):
        raise ValueError("Cannot chart incomplete CRC/USD cashflow bands")

    for offset, window in enumerate(projection.windows, 5):
        cells: list[object] = [window.label]
        for currency in ("CRC", "USD"):
            row = indexed[currency, window.key]
            # A missing payment schedule is not a certified zero-flow interval.
            cells.append(row.net_flow if row.event_count else "N/D")
            cells.append(row.accumulated_flow if projection.source_flow_count else "N/D")
        for column, value in enumerate(cells, 1):
            cell = sheet.cell(row=offset, column=column, value=value)
            if column > 1 and value != "N/D":
                cell.number_format = "#,##0.00;[Red](#,##0.00)"
    if not projection.source_flow_count:
        sheet["A24"] = "Sin eventos de caja aprobados: no se generan gráficos."
        return

    categories = Reference(sheet, min_col=1, min_row=5, max_row=sheet.max_row)
    for currency, net_column, gap_column, position in (
        ("CRC", 2, 3, "G4"),
        ("USD", 4, 5, "G24"),
    ):
        if not any(row.event_count for row in projection.rows if row.currency == currency):
            continue
        net_chart = BarChart()
        net_chart.type = "col"
        net_chart.style = 10
        net_chart.title = f"{currency} · flujo neto documentado (parcial)"
        net_chart.y_axis.title = f"Moneda original: {currency}"
        net_chart.x_axis.title = "Banda de proyección"
        net_chart.add_data(
            Reference(sheet, min_col=net_column, min_row=4, max_row=sheet.max_row),
            titles_from_data=True,
        )
        net_chart.set_categories(categories)
        net_chart.width = 19
        net_chart.height = 9
        net_chart.legend = None
        sheet.add_chart(net_chart, position)

        gap_chart = LineChart()
        gap_chart.style = 13
        gap_chart.title = f"{currency} · GAP acumulado documentado (parcial)"
        gap_chart.y_axis.title = f"Moneda original: {currency}"
        gap_chart.x_axis.title = "Banda de proyección"
        gap_chart.add_data(
            Reference(sheet, min_col=gap_column, min_row=4, max_row=sheet.max_row),
            titles_from_data=True,
        )
        gap_chart.set_categories(categories)
        gap_chart.width = 19
        gap_chart.height = 9
        gap_chart.legend = None
        sheet.add_chart(gap_chart, "P4" if currency == "CRC" else "P24")
