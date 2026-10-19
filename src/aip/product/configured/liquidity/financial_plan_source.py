"""Read the approved-format financial plan as separate, unallocated liquidity evidence.

The source contains CRC-million *equivalent* budgets, not certified dated
currency cash flows. Read-only: no filename is a presumed approval. Amounts
cannot enter the SUGEF 2-10 waterfall without date, native currency and
non-duplication controls in financial_plan_cashflow_bridge.
"""

from __future__ import annotations

import hashlib
from calendar import monthrange
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# The recognized plan drivers are consciously limited to the monthly calendar.
# Annual profit and expense budgets must NOT be divided by twelve nor added on
# top of already-scheduled contractual interest or operating payments.
_DIRECT_MONTHLY = {
    "Colocación bruta": ("ASSET", "OUTFLOW"),
    "Emisión nueva CAP (promedio)": ("LIABILITY", "INFLOW"),
    "Aportes capital social": ("BUSINESS_EVENT", "INFLOW"),
    "Renuncias capital social": ("BUSINESS_EVENT", "OUTFLOW"),
}
_NONCASH = frozenset(
    {
        "Gasto estimación cartera",
        "Traslado a incobrables",
        "Capitalización excedentes",
    }
)
_PARAMETER_ONLY = frozenset({"Tasa EML"})
_MONTHS = (
    "ene-26", "feb-26", "mar-26", "abr-26", "may-26", "jun-26",
    "jul-26", "ago-26", "sept-26", "oct-26", "nov-26", "dic-26",
)

# Approved business-plan source IDs are classifications, NOT model approvals.
_PLAN_LIQUIDITY_ROLES = {
    1: "SCENARIO_FX_REFERENCE_NOT_OFFICIAL_SPOT",
    2: "RATE_DRIVER_NOT_CASH",
    3: "RATE_DRIVER_NOT_CASH",
    4: "RATE_DRIVER_NOT_CASH",
    5: "NEW_LENDING_RATE_PASS_THROUGH",
    6: "NEW_DEPOSIT_RATE_PASS_THROUGH",
    7: "INTEREST_BUDGET_NOT_CONTRACTUAL_CASH",
    9: "INVESTMENT_INTEREST_BUDGET_AVOID_MASTER_DUPLICATION",
    10: "REINVESTMENT_YIELD_REQUIRES_APPROVED_TRADE_SCENARIO",
    14: "DEPOSIT_INTEREST_BUDGET_AVOID_CONTRACT_DUPLICATION",
    15: "CRC_NEW_TERM_DEPOSIT_RATE_DRIVER",
    16: "USD_NEW_TERM_DEPOSIT_RATE_DRIVER",
    17: "BORROWING_INTEREST_BUDGET_AVOID_SCHEDULE_DUPLICATION",
    25: "INCLUDED_IN_OPERATING_EXPENSES_DUPLICATE_GUARD",
    26: "INCLUDED_IN_OPERATING_EXPENSES_DUPLICATE_GUARD",
    30: "NONCASH_PROVISION_EXCLUDED",
    31: "NONCASH_PROVISION_EXCLUDED",
    34: "GROSS_CREDIT_DISBURSEMENT_PLANNING",
    35: "USD_LOAN_SHARE_NOT_CERTIFIED_CURRENCY_SPLIT",
    37: "NONCASH_WRITE_OFF_EXCLUDED",
    38: "IMPLICIT_PRINCIPAL_RECOVERY_NOT_CONTRACTUAL",
    40: "INVESTMENT_MATURITY_BUDGET_AVOID_MASTER_DUPLICATION",
    42: "EML_RATE_RESTRICTION_REQUIRES_BASE_BALANCES",
    43: "EML_RATE_RESTRICTION_REQUIRES_BASE_BALANCES",
    44: "EML_REQUIRED_RESERVE_NOT_SPENDABLE_CASH",
    49: "NEW_TERM_DEPOSIT_ISSUANCE_AVOID_RENEWAL_DUPLICATION",
    50: "DEPOSIT_RENEWAL_ASSUMPTION_PENDING_APPROVAL",
    52: "BORROWING_AMORTIZATION_AVOID_SCHEDULE_DUPLICATION",
    55: "NONCASH_EXCESS_CAPITALIZATION_EXCLUDED",
}


@dataclass(frozen=True, slots=True)
class PlanParameter:
    plan_id: int
    name: str
    value: str | None
    unit: str
    periodicity: str
    cashflow_type: str
    plan_reference: str
    source_cell: str
    observation: str
    liquidity_role: str = "PLAN_REFERENCE_NOT_DATED_CASH"


@dataclass(frozen=True, slots=True)
class MonthlyFinancialPlanInput:
    driver: str
    year: int
    month: int
    reporting_currency: str
    reported_value: Decimal
    source_cell: str
    source_sha256: str
    area: str | None
    side: str | None
    classification: str
    status: str

    @property
    def period_start(self) -> date:
        return date(self.year, self.month, 1)

    @property
    def period_end(self) -> date:
        return date(self.year, self.month, monthrange(self.year, self.month)[1])


@dataclass(frozen=True, slots=True)
class FinancialPlanSourceAudit:
    cutoff: date
    source_name: str
    sha256: str
    plan_year: int
    parameters: tuple[PlanParameter, ...]
    monthly_inputs: tuple[MonthlyFinancialPlanInput, ...]
    observations: tuple[str, ...]
    status: str = "PLANNING_SOURCE_EVIDENCE_ONLY_APPROVAL_AND_TIMING_REQUIRED"

    @property
    def eligible_cashflows(self) -> tuple[()]:
        """Never expose source forecast rows as approved payment events."""
        return ()


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        amount = Decimal(str(value).strip())
    except (ValueError, InvalidOperation):
        return None
    return amount if amount.is_finite() else None


def read_financial_plan_assumptions(
    source_file: str | Path, *, cutoff: date
) -> FinancialPlanSourceAudit:
    """Fail closed on an unfamiliar plan format, calendar or inconsistent totals."""
    source = Path(source_file)
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise ValueError("Expected existing .xlsx financial planning matrix")
    if source.stat().st_size > 20_000_000:
        raise ValueError("Financial planning workbook exceeds safe read limit")
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    source_sha = digest.hexdigest()

    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        if not {"Matriz Supuestos", "Calendario mensual"}.issubset(workbook.sheetnames):
            raise ValueError("Plan workbook is missing the institutional source sheets")
        matrix = workbook["Matriz Supuestos"]
        calendar = workbook["Calendario mensual"]
        headers = tuple(calendar.cell(3, col).value for col in range(1, 15))
        if headers[0] != "Driver" or tuple(headers[1:13]) != _MONTHS:
            raise ValueError("Plan calendar months differ from certified source layout")
        if headers[13] != "Total 2026":
            raise ValueError("Plan annual total column is missing")
        if tuple(matrix.cell(4, col).value for col in range(1, 5)) != (
            "ID", "Actividad", "Rubro", "Supuesto / Driver"
        ):
            raise ValueError("Plan matrix headers do not match configured layout")

        parameters: list[PlanParameter] = []
        for row in range(5, matrix.max_row + 1):
            ident = matrix.cell(row, 1).value
            if not isinstance(ident, int) or isinstance(ident, bool):
                continue
            label = str(matrix.cell(row, 4).value or "").strip()
            raw = matrix.cell(row, 5).value
            if raw is None:
                value = None
            elif isinstance(raw, (int, float, str)) and not isinstance(raw, bool):
                value = str(raw)
            else:
                value = None
            parameters.append(
                PlanParameter(
                    plan_id=ident,
                    name=label,
                    value=value,
                    unit=str(matrix.cell(row, 6).value or ""),
                    periodicity=str(matrix.cell(row, 7).value or ""),
                    cashflow_type=str(matrix.cell(row, 8).value or ""),
                    plan_reference=str(matrix.cell(row, 11).value or ""),
                    source_cell=f"'Matriz Supuestos'!E{row}",
                    observation=str(matrix.cell(row, 12).value or ""),
                    liquidity_role=_PLAN_LIQUIDITY_ROLES.get(
                        ident, "PLAN_REFERENCE_NOT_DATED_CASH"
                    ),
                )
            )

        monthly: list[MonthlyFinancialPlanInput] = []
        seen: set[str] = set()
        warnings: list[str] = []
        known_drivers = frozenset(_DIRECT_MONTHLY) | _NONCASH | _PARAMETER_ONLY
        for row in range(4, calendar.max_row + 1):
            driver = str(calendar.cell(row, 1).value or "").strip()
            if driver not in known_drivers:
                continue
            if driver in seen:
                raise ValueError("Duplicate source calendar driver: " + driver)
            seen.add(driver)
            area, side = _DIRECT_MONTHLY.get(driver, (None, None))
            if driver in _NONCASH:
                classification = "NONCASH_EXCLUDED"
            elif driver in _PARAMETER_ONLY:
                classification = "PARAMETER_NOT_CASHFLOW"
            else:
                classification = "PLANNED_POTENTIAL_CASHFLOW"
            monthly_values: list[Decimal] = []
            for month in range(1, 13):
                raw = calendar.cell(row, month + 1).value
                value = _decimal(raw)
                if value is None:
                    raise ValueError(f"Missing or invalid plan driver: {driver} month {month}")
                if value < 0:
                    raise ValueError("Monthly budget inputs cannot be negative")
                monthly_values.append(value)
                status = classification
                if area is not None:
                    if date(2026, month, 1) <= date(cutoff.year, cutoff.month, 1):
                        status = "HISTORICAL_OR_PARTIAL_MONTH_NOT_PROJECTED"
                    else:
                        status = "PENDING_APPROVAL_CURRENCY_DATE_AND_DUPLICATION_REVIEW"
                monthly.append(
                    MonthlyFinancialPlanInput(
                        driver=driver,
                        year=2026,
                        month=month,
                        reporting_currency="CRC_MILLION_EQUIVALENT",
                        reported_value=value,
                        source_cell=f"'Calendario mensual'!{chr(ord('B') + month - 1)}{row}",
                        source_sha256=source_sha,
                        area=area,
                        side=side,
                        classification=classification,
                        status=status,
                    )
                )
            annual = _decimal(calendar.cell(row, 14).value)
            if driver in _DIRECT_MONTHLY or driver in _NONCASH:
                if annual is None or annual != sum(monthly_values, Decimal(0)):
                    raise ValueError("Plan monthly calendar does not reconcile annual total: " + driver)
            elif annual is not None:
                warnings.append("Parameter row contains an unexpected annual cash total")
        if set(_DIRECT_MONTHLY) - seen:
            raise ValueError("Plan calendar is missing required projected cash drivers")
        if cutoff.year != 2026:
            warnings.append("PLAN_YEAR_DIFFERS_FROM_REQUESTED_CUTOFF_NO_ROLLOVER")
        warnings.append("PLAN_DOES_NOT_PROVE_ALCO_APPROVAL_OR_DATED_CASH")
        warnings.append("MONTHLY_CRC_EQUIVALENT_CANNOT_BE_SPLIT_INTO_CRC_USD_WITHOUT_FX_EVIDENCE")
        return FinancialPlanSourceAudit(
            cutoff=cutoff,
            source_name=source.name,
            sha256=source_sha,
            plan_year=2026,
            parameters=tuple(parameters),
            monthly_inputs=tuple(monthly),
            observations=tuple(warnings),
        )
    finally:
        workbook.close()
