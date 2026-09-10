from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from aip.product.configured.irrbb.physical_source_registry import BORROWING_WORKBOOK_SOURCE


@dataclass(frozen=True, slots=True)
class BorrowingWorkbookSheetTopology:
    """Observed worksheet topology without assigning business meaning to cells."""

    sheet_name: str
    visibility: str
    worksheet_max_row: int
    worksheet_max_column: int
    observed_non_empty_row_count: int
    observed_non_empty_cell_count: int
    observed_last_non_empty_row: int
    observed_last_non_empty_column: int


@dataclass(frozen=True, slots=True)
class BorrowingWorkbookDiscovery:
    """Auditable identity and topology of one governed borrowing workbook."""

    source_id: str
    source_reference: str
    source_file_name: str
    file_sha256: str
    sheets: tuple[BorrowingWorkbookSheetTopology, ...]


@dataclass(frozen=True, slots=True)
class BorrowingWorkbookHeaderCellEvidence:
    """Exact text observed in one declared header cell."""

    column_index: int
    column_letter: str
    label: str | None


@dataclass(frozen=True, slots=True)
class BorrowingWorkbookHeaderInspection:
    """Evidence for an explicitly selected worksheet/header row."""

    discovery: BorrowingWorkbookDiscovery
    sheet_name: str
    header_row: int
    cells: tuple[BorrowingWorkbookHeaderCellEvidence, ...]
    blank_column_indexes: tuple[int, ...]
    duplicate_labels: tuple[str, ...]


class BorrowingWorkbookSchemaInspector:
    """Inspect the governed obligations workbook without inferring its RTILB schema.

    Discovery records workbook identity and worksheet topology. Header inspection
    requires the caller to declare both worksheet and header row; no heuristic is
    allowed to select them. The inspector never maps a row to a canonical RTILB
    position and never upgrades source certification.
    """

    _SUPPORTED_SUFFIX = ".xlsx"
    _HASH_CHUNK_SIZE = 1024 * 1024

    def discover(self, path: str | Path) -> BorrowingWorkbookDiscovery:
        file_path = self._validate_source_path(path)
        digest = self._sha256(file_path)
        workbook = self._load_workbook(file_path)
        try:
            sheets = tuple(self._inspect_sheet_topology(sheet) for sheet in workbook.worksheets)
        finally:
            workbook.close()
        self._require_unchanged_file(file_path, expected_digest=digest)

        if not sheets:
            raise ValueError("borrowing workbook must contain at least one worksheet")

        return BorrowingWorkbookDiscovery(
            source_id=BORROWING_WORKBOOK_SOURCE.source_id,
            source_reference=self._source_reference(file_path.name, digest),
            source_file_name=file_path.name,
            file_sha256=digest,
            sheets=sheets,
        )

    def inspect_declared_header(
        self,
        path: str | Path,
        *,
        sheet_name: str,
        header_row: int,
    ) -> BorrowingWorkbookHeaderInspection:
        if not sheet_name.strip():
            raise ValueError("borrowing workbook sheet_name is required")
        if header_row <= 0:
            raise ValueError("borrowing workbook header_row must be positive")

        discovery = self.discover(path)
        topology = self._require_sheet(discovery, sheet_name)
        if topology.observed_last_non_empty_row == 0:
            raise ValueError(f"borrowing workbook sheet {sheet_name!r} contains no observed values")
        if header_row > topology.observed_last_non_empty_row:
            raise ValueError(
                "borrowing workbook header_row exceeds the last observed non-empty row "
                f"for sheet {sheet_name!r}"
            )

        file_path = self._validate_source_path(path)
        self._require_unchanged_file(file_path, expected_digest=discovery.file_sha256)
        workbook = self._load_workbook(file_path)
        try:
            sheet = workbook[sheet_name]
            cells = self._read_declared_header(
                sheet,
                header_row=header_row,
                max_column=topology.observed_last_non_empty_column,
            )
        finally:
            workbook.close()
        self._require_unchanged_file(file_path, expected_digest=discovery.file_sha256)

        blank_column_indexes = tuple(cell.column_index for cell in cells if cell.label is None)
        duplicate_labels = self._duplicate_labels(cells)
        if all(cell.label is None for cell in cells):
            raise ValueError("borrowing workbook declared header row contains no text labels")

        return BorrowingWorkbookHeaderInspection(
            discovery=discovery,
            sheet_name=sheet_name,
            header_row=header_row,
            cells=cells,
            blank_column_indexes=blank_column_indexes,
            duplicate_labels=duplicate_labels,
        )

    @classmethod
    def _validate_source_path(cls, path: str | Path) -> Path:
        file_path = Path(path)
        if not file_path.exists():
            raise FileNotFoundError(f"borrowing workbook was not found: {file_path.name}")
        if not file_path.is_file():
            raise ValueError("borrowing workbook source must be a regular file")
        if file_path.suffix.casefold() != cls._SUPPORTED_SUFFIX:
            raise ValueError("borrowing workbook source must use the .xlsx format")
        if file_path.name != BORROWING_WORKBOOK_SOURCE.logical_name:
            raise ValueError(
                "borrowing workbook filename does not match the governed physical source identity"
            )
        return file_path

    @staticmethod
    def _load_workbook(path: Path) -> Any:
        try:
            return openpyxl.load_workbook(path, read_only=True, data_only=False)
        except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as exc:
            raise ValueError(f"borrowing workbook could not be read: {path.name}") from exc

    @classmethod
    def _sha256(cls, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            while chunk := stream.read(cls._HASH_CHUNK_SIZE):
                digest.update(chunk)
        return digest.hexdigest()

    @classmethod
    def _require_unchanged_file(cls, path: Path, *, expected_digest: str) -> None:
        if cls._sha256(path) != expected_digest:
            raise ValueError("borrowing workbook changed while inspection evidence was being captured")

    @staticmethod
    def _source_reference(file_name: str, digest: str) -> str:
        return f"{BORROWING_WORKBOOK_SOURCE.source_id}:{file_name}#sha256={digest}"

    @classmethod
    def _inspect_sheet_topology(cls, sheet: Any) -> BorrowingWorkbookSheetTopology:
        observed_non_empty_row_count = 0
        observed_non_empty_cell_count = 0
        observed_last_non_empty_row = 0
        observed_last_non_empty_column = 0

        for row in sheet.iter_rows():
            row_has_value = False
            for cell in row:
                if not cls._has_observed_value(cell.value):
                    continue
                row_has_value = True
                observed_non_empty_cell_count += 1
                observed_last_non_empty_row = max(observed_last_non_empty_row, cell.row)
                observed_last_non_empty_column = max(
                    observed_last_non_empty_column,
                    cell.column,
                )
            if row_has_value:
                observed_non_empty_row_count += 1

        return BorrowingWorkbookSheetTopology(
            sheet_name=sheet.title,
            visibility=sheet.sheet_state,
            worksheet_max_row=int(sheet.max_row or 0),
            worksheet_max_column=int(sheet.max_column or 0),
            observed_non_empty_row_count=observed_non_empty_row_count,
            observed_non_empty_cell_count=observed_non_empty_cell_count,
            observed_last_non_empty_row=observed_last_non_empty_row,
            observed_last_non_empty_column=observed_last_non_empty_column,
        )

    @staticmethod
    def _has_observed_value(value: object) -> bool:
        if value is None:
            return False
        if isinstance(value, str) and not value.strip():
            return False
        return True

    @staticmethod
    def _require_sheet(
        discovery: BorrowingWorkbookDiscovery,
        sheet_name: str,
    ) -> BorrowingWorkbookSheetTopology:
        for topology in discovery.sheets:
            if topology.sheet_name == sheet_name:
                return topology
        raise KeyError(f"borrowing workbook worksheet was not found: {sheet_name}")

    @staticmethod
    def _read_declared_header(
        sheet: Any,
        *,
        header_row: int,
        max_column: int,
    ) -> tuple[BorrowingWorkbookHeaderCellEvidence, ...]:
        row = next(
            sheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                min_col=1,
                max_col=max_column,
            )
        )
        evidence: list[BorrowingWorkbookHeaderCellEvidence] = []
        for cell in row:
            value = cell.value
            if cell.data_type == "f":
                raise ValueError(
                    "borrowing workbook declared header cells cannot be formulas; "
                    f"formula observed at {cell.coordinate}"
                )
            if value is None or (isinstance(value, str) and not value.strip()):
                label = None
            elif not isinstance(value, str):
                raise ValueError(
                    "borrowing workbook declared header cells must contain text or be blank; "
                    f"non-text value observed at {cell.coordinate}"
                )
            else:
                label = value
            evidence.append(
                BorrowingWorkbookHeaderCellEvidence(
                    column_index=cell.column,
                    column_letter=get_column_letter(cell.column),
                    label=label,
                )
            )
        return tuple(evidence)

    @staticmethod
    def _duplicate_labels(
        cells: tuple[BorrowingWorkbookHeaderCellEvidence, ...],
    ) -> tuple[str, ...]:
        first_exact_label_by_key: dict[str, str] = {}
        duplicate_keys: set[str] = set()
        for cell in cells:
            if cell.label is None:
                continue
            key = " ".join(cell.label.split()).casefold()
            if key in first_exact_label_by_key:
                duplicate_keys.add(key)
            else:
                first_exact_label_by_key[key] = cell.label
        return tuple(first_exact_label_by_key[key] for key in sorted(duplicate_keys))
